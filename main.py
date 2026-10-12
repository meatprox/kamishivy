from __future__ import annotations

import configparser
import csv
import io
import threading
import tkinter as tk
import tkinter.font as tkfont
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from tkinter import colorchooser, filedialog, messagebox, ttk
from typing import Any

from app.channels import (
    DEFAULT_ENCODING,
    DEFAULT_RAW_LENGTH,
    DEFAULT_RECV_RETRIES,
    DEFAULT_RECV_WAIT,
    ENCODING_CHOICES,
    MAX_CHANNELS,
    RAW_CHANNEL,
    RAW_CHANNEL_NAME,
    ChannelHub,
    channel_label,
    clamp_raw_length,
    clamp_recv_retries,
    clamp_recv_wait,
    encoding_label,
    list_com_ports,
    resolve_channel,
)
from app.design import (
    DEFAULTS,
    default_design,
    font_tuple,
    hex_to_rgb,
    load_design,
    load_design_preset,
    rgb_hex,
    save_design,
    save_design_preset,
)
from app.logger import SessionLogger, heading_cells, is_heading_row, normalize_headings, read_log_records
from app.parse import (
    PARSE_MODE_CSV,
    PARSE_MODE_LABELS,
    PARSE_MODE_NONE,
    PARSE_MODE_SLICE,
    PARSE_MODES,
    PARSE_OUTPUT_LABELS,
    PARSE_OUTPUT_TEXT,
    PARSE_OUTPUTS,
    process_received,
)
from app.paths import app_root
from ui.base import BaseUI, UIContext, create_ui, preload_uis, strip_trailing_cui_empties


ROOT = app_root()
CONFIG_PATH = ROOT / "config.ini"
PORT_INI_PATH = ROOT / "port.ini"
LEGACY_COM_MAP_PATH = ROOT / "config.txt"
CUI_CANDIDATES = ("cui.xlsx", "cui.xlsm", "cui.xls", "cui.csv")
UNUSED_PORT = "未使用"
BAUD_CHOICES = (
    "1200",
    "2400",
    "4800",
    "9600",
    "19200",
    "38400",
    "57600",
    "115200",
    "230400",
    "460800",
    "921600",
)
RAW_LENGTH_CHOICES = ("1", "2", "4", "8", "16", "32", "64", "128", "256")

LAMP_OFF = "#8a8a8a"
LAMP_ON = "#2ecc71"
LAMP_ERROR = "#e74c3c"
LAMP_WAIT = "#f1c40f"
NEXT_SIZE_MODE_DEFAULT = "percent"
NEXT_WIDTH_DEFAULT = 20
NEXT_HEIGHT_DEFAULT = 7


def load_config(path: Path | None = None) -> configparser.ConfigParser:
    parser = configparser.ConfigParser(interpolation=None)
    parser.optionxform = str
    parser.read(path or CONFIG_PATH, encoding="utf-8")
    return parser


def save_config(parser: configparser.ConfigParser) -> None:
    limit = channel_max(parser)
    if not parser.has_section("com"):
        parser.add_section("com")
    if not parser.has_section("channels"):
        parser.add_section("channels")
    use = get(parser, "com", "use", "1") or "1"
    lamps = get(parser, "com", "lamp", "")
    lines = [
        "; kamishivy 利用者向け設定",
        "; このファイルはソフトウェア使用者が変更する項目です。通信設定もここに保存し、次回起動へ持ち越します。",
        "",
        "[window]",
        f"title = {get(parser, 'window', 'title', 'kamishivy') or 'kamishivy'}",
        f"width = {get_int(parser, 'window', 'width', 960)}",
        f"height = {get_int(parser, 'window', 'height', 640)}",
        "",
        "[display]",
        "; fullscreen: 1なら全画面。hide_close: 1なら×ボタンを隠す（全画面のときだけ）",
        f"fullscreen = {get_int(parser, 'display', 'fullscreen', 0)}",
        f"hide_close = {get_int(parser, 'display', 'hide_close', 0)}",
        "; show_page: 1ならメニュー左に 現在/総ページ を出す",
        f"show_page = {get_int(parser, 'display', 'show_page', 1)}",
        "; next_size_mode: percent=窓全体に対する割合（横=幅、縦=高さ） / px=絶対サイズ",
        f"next_size_mode = {get(parser, 'display', 'next_size_mode', NEXT_SIZE_MODE_DEFAULT) or NEXT_SIZE_MODE_DEFAULT}",
        f"next_width = {get(parser, 'display', 'next_width', str(NEXT_WIDTH_DEFAULT)) or str(NEXT_WIDTH_DEFAULT)}",
        f"next_height = {get(parser, 'display', 'next_height', str(NEXT_HEIGHT_DEFAULT)) or str(NEXT_HEIGHT_DEFAULT)}",
        "; next_align: left / center / right",
        f"next_align = {get(parser, 'display', 'next_align', 'center') or 'center'}",
        "",
        "[serial]",
        f"baudrate = {get_int(parser, 'serial', 'baudrate', 9600)}",
        f"bytesize = {get_int(parser, 'serial', 'bytesize', 8)}",
        f"parity = {get(parser, 'serial', 'parity', 'N') or 'N'}",
        f"stopbits = {get(parser, 'serial', 'stopbits', '1') or '1'}",
        "; timeout: 1回の受信待ち（秒）。0 なら待たずに次を見る（リアルタイム向き）",
        f"timeout = {clamp_recv_wait(get(parser, 'serial', 'timeout', str(DEFAULT_RECV_WAIT)))}",
        "; retries: 取れなかったときの再試行回数。0 なら再試行しない",
        f"retries = {clamp_recv_retries(get(parser, 'serial', 'retries', str(DEFAULT_RECV_RETRIES)))}",
        "",
        "[com]",
        "; 1ならCOMを使う（接続ボタン表示）。0なら使わない（接続ボタンを隠す）",
        f"use = {use}",
        "; 接続ランプを出すCOMポート（カンマ区切り。番号は使わない）",
        f"lamp = {lamps}",
        "",
        "[channels]",
        "; データ番号の上限と、番号に割り当てる COM ポート。空なら未使用。",
        f"; {RAW_CHANNEL_NAME}は区切りなし（固定長）読み込み用。COM は人が選ぶ。",
        f"max = {limit}",
    ]
    for number in range(1, limit + 1):
        if number == RAW_CHANNEL:
            continue
        port = _channel_assignment(parser, number)
        lines.append(f"{number} = {port}")
    lines.append(f"{RAW_CHANNEL_NAME} = {_channel_assignment(parser, RAW_CHANNEL)}")
    lines.append(f"; {RAW_CHANNEL_NAME}が1回に読むバイト数。")
    lines.append(f"raw_length = {load_raw_length(parser)}")
    lines.extend(
        [
            "",
            "[baud]",
            "; ポートごとのボーレート。未記入のポートは [serial] baudrate を使う。",
        ]
    )
    if parser.has_section("baud"):
        for key in parser.options("baud"):
            port = key.strip()
            if not port:
                continue
            rate = parser.get("baud", key, fallback="").strip()
            if rate:
                lines.append(f"{port} = {rate}")
    lines.extend(
        [
            "",
            "[encoding]",
            "; ポートごとの文字コード。未記入のポートは UTF-8。",
        ]
    )
    if parser.has_section("encoding"):
        for key in parser.options("encoding"):
            port = key.strip()
            if not port:
                continue
            value = parser.get("encoding", key, fallback="").strip()
            if value:
                lines.append(f"{port} = {encoding_label(value)}")
    lines.extend(
        [
            "",
            "[timeout]",
            "; ポートごとの受信待ち（秒）。未記入のポートは [serial] timeout を使う。",
        ]
    )
    if parser.has_section("timeout"):
        for key in parser.options("timeout"):
            port = key.strip()
            if not port:
                continue
            value = parser.get("timeout", key, fallback="").strip()
            if value:
                lines.append(f"{port} = {clamp_recv_wait(value):g}")
    lines.extend(
        [
            "",
            "[retries]",
            "; ポートごとの再試行回数。未記入のポートは [serial] retries を使う。",
        ]
    )
    if parser.has_section("retries"):
        for key in parser.options("retries"):
            port = key.strip()
            if not port:
                continue
            value = parser.get("retries", key, fallback="").strip()
            if value:
                lines.append(f"{port} = {clamp_recv_retries(value)}")
    lines.extend(
        [
            "",
            "[parse]",
            "; mode: none=加工しない / slice=文字列の一部 / csv=カンマ区切りの何番目",
            "; output: text=文字 / number=数値",
            "; start,end は slice の1始まり位置。csv は何番目（1始まり）。",
        ]
    )
    parse_limit = max(limit, RAW_CHANNEL)
    for number in range(1, parse_limit + 1):
        spec = load_parse_spec(parser, number)
        key = _parse_channel_key(number)
        lines.append(f"{key}_mode = {spec['mode']}")
        lines.append(f"{key}_output = {spec['output']}")
        lines.append(f"{key}_start = {spec['start']}")
        lines.append(f"{key}_end = {spec['end']}")
        lines.append(f"{key}_csv = {spec['csv']}")
    lines.extend(
        [
            "",
            "[cui]",
            "; Excel（.xlsx）または CSV。両方ある場合はここの指定が優先。",
            "; remember: 1なら前回のCUIを使う。0なら起動時に読み込み窓を出す。",
            f"file = {get(parser, 'cui', 'file', '')}",
            f"remember = {get_int(parser, 'cui', 'remember', 1)}",
            "",
            "[log]",
            f"directory = {get(parser, 'log', 'directory', 'logs') or 'logs'}",
            f"backup = {get(parser, 'log', 'backup', '')}",
            "",
            "[sound]",
            "; next: 1なら次へを押したときに sound1.wav を鳴らす。0なら鳴らさない。",
            f"next = {get_int(parser, 'sound', 'next', 1)}",
            "",
            "[nav]",
            "; back: 1ならメニューの戻るを使う。0なら使えない。",
            f"back = {get_int(parser, 'nav', 'back', 1)}",
            "; comment: 1ならメニューの次へでコメントを書く。0なら入力窓を出さない。",
            f"comment = {get_int(parser, 'nav', 'comment', 0)}",
            "",
        ]
    )
    CONFIG_PATH.write_text("\n".join(lines), encoding="utf-8")


def get(parser: configparser.ConfigParser, section: str, key: str, fallback: Any = "") -> str:
    if not parser.has_section(section):
        return str(fallback).strip()
    return parser.get(section, key, fallback=str(fallback)).strip()


def get_int(parser: configparser.ConfigParser, section: str, key: str, fallback: int) -> int:
    if not parser.has_section(section):
        return fallback
    try:
        return parser.getint(section, key, fallback=fallback)
    except (ValueError, TypeError):
        return fallback


def _parse_channel_key(number: int) -> str:
    return RAW_CHANNEL_NAME if number == RAW_CHANNEL else str(number)


def load_parse_spec(parser: configparser.ConfigParser, number: int) -> dict[str, Any]:
    key = _parse_channel_key(number)
    mode = (get(parser, "parse", f"{key}_mode", PARSE_MODE_NONE) or PARSE_MODE_NONE).lower()
    if mode not in PARSE_MODES:
        mode = PARSE_MODE_NONE
    output = (get(parser, "parse", f"{key}_output", PARSE_OUTPUT_TEXT) or PARSE_OUTPUT_TEXT).lower()
    if output not in PARSE_OUTPUTS:
        output = PARSE_OUTPUT_TEXT
    start = max(1, min(4096, get_int(parser, "parse", f"{key}_start", 1)))
    end = max(1, min(4096, get_int(parser, "parse", f"{key}_end", 1)))
    csv_n = max(1, min(99, get_int(parser, "parse", f"{key}_csv", 1)))
    return {"mode": mode, "output": output, "start": start, "end": end, "csv": csv_n}


def get_float(parser: configparser.ConfigParser, section: str, key: str, fallback: float) -> float:
    if not parser.has_section(section):
        return fallback
    try:
        return float(parser.get(section, key, fallback=str(fallback)))
    except (ValueError, TypeError):
        return fallback


def channel_max(parser: configparser.ConfigParser) -> int:
    value = get_int(parser, "channels", "max", MAX_CHANNELS)
    return min(MAX_CHANNELS, max(1, value))


def _parse_com_map_text(text: str, limit: int) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or line.startswith(";"):
            continue
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        if key.strip().lower() in ("max", "raw_length"):
            continue
        number = resolve_channel(key.strip())
        if number < 1:
            continue
        port = value.strip()
        if 1 <= number <= limit and port:
            mapping[number] = port
    return mapping


def _channel_assignment(parser: configparser.ConfigParser, number: int) -> str:
    if not parser.has_section("channels"):
        return ""
    keys = [str(number)]
    if number == RAW_CHANNEL:
        keys = [RAW_CHANNEL_NAME, "A", "a", str(RAW_CHANNEL)]
    for key in keys:
        if parser.has_option("channels", key):
            return parser.get("channels", key, fallback="").strip()
    return ""


def _channel_keys_present(parser: configparser.ConfigParser) -> bool:
    if not parser.has_section("channels"):
        return False
    return any(
        key.strip().isdigit() or resolve_channel(key) >= 1
        for key in parser.options("channels")
        if key.strip().lower() not in ("max", "raw_length")
    )


def _ini_channel_ports(parser: configparser.ConfigParser) -> dict[int, str]:
    limit = channel_max(parser)
    mapping: dict[int, str] = {}
    if parser.has_section("channels"):
        for key, value in parser.items("channels"):
            if key.strip().lower() in ("max", "raw_length"):
                continue
            number = resolve_channel(key)
            port = (value or "").strip()
            if 1 <= number <= limit and port:
                mapping[number] = port
    if mapping or _channel_keys_present(parser):
        return mapping
    port = get(parser, "serial", "port", "")
    if port:
        mapping[1] = port
    return mapping


def _read_port_ini(path: Path, limit: int) -> dict[int, str] | None:
    try:
        return _parse_com_map_text(path.read_text(encoding="utf-8"), limit)
    except OSError:
        return None


def load_channel_ports(parser: configparser.ConfigParser) -> dict[int, str]:
    if _channel_keys_present(parser) or parser.has_section("com"):
        return _ini_channel_ports(parser)
    limit = channel_max(parser)
    if PORT_INI_PATH.exists():
        mapping = _read_port_ini(PORT_INI_PATH, limit)
        if mapping is not None:
            return mapping
    if LEGACY_COM_MAP_PATH.exists():
        mapping = _read_port_ini(LEGACY_COM_MAP_PATH, limit)
        if mapping is not None:
            return mapping
    return _ini_channel_ports(parser)


def save_channel_ports(
    mapping: dict[int, str],
    limit: int,
    lamps: list[str] | None = None,
    use_com: bool | None = None,
    baudrates: dict[str, int] | None = None,
    encodings: dict[str, str] | None = None,
    timeouts: dict[str, float] | None = None,
    retries: dict[str, int] | None = None,
    raw_length: int | None = None,
    parser: configparser.ConfigParser | None = None,
) -> configparser.ConfigParser:
    parser = parser or load_config()
    if lamps is None:
        lamps = load_lamp_ports(parser)
    if use_com is None:
        use_com = load_com_enabled(parser)
    if baudrates is None:
        baudrates = load_port_baudrates(parser)
    if encodings is None:
        encodings = load_port_encodings(parser)
    if timeouts is None:
        timeouts = load_port_timeouts(parser)
    if retries is None:
        retries = load_port_retries(parser)
    if raw_length is None:
        raw_length = load_raw_length(parser)
    if not parser.has_section("com"):
        parser.add_section("com")
    if not parser.has_section("channels"):
        parser.add_section("channels")
    parser.set("com", "use", "1" if use_com else "0")
    parser.set("com", "lamp", ",".join(_unique_port_names(lamps)))
    parser.set("channels", "max", str(limit))
    for number in range(1, limit + 1):
        if number == RAW_CHANNEL:
            continue
        parser.set("channels", str(number), mapping.get(number, ""))
    parser.set("channels", RAW_CHANNEL_NAME, mapping.get(RAW_CHANNEL, ""))
    for legacy in ("10", "A", "a"):
        if legacy != RAW_CHANNEL_NAME and parser.has_option("channels", legacy):
            parser.remove_option("channels", legacy)
    parser.set("channels", "raw_length", str(clamp_raw_length(raw_length)))
    if parser.has_section("baud"):
        parser.remove_section("baud")
    parser.add_section("baud")
    for port, rate in sorted(baudrates.items(), key=lambda item: item[0].lower()):
        name = (port or "").strip()
        if not name:
            continue
        parser.set("baud", name, str(int(rate)))
    if parser.has_section("encoding"):
        parser.remove_section("encoding")
    parser.add_section("encoding")
    for port, label in sorted(encodings.items(), key=lambda item: item[0].lower()):
        name = (port or "").strip()
        if not name:
            continue
        parser.set("encoding", name, encoding_label(label))
    if parser.has_section("timeout"):
        parser.remove_section("timeout")
    parser.add_section("timeout")
    for port, wait in sorted(timeouts.items(), key=lambda item: item[0].lower()):
        name = (port or "").strip()
        if not name:
            continue
        parser.set("timeout", name, f"{clamp_recv_wait(wait):g}")
    if parser.has_section("retries"):
        parser.remove_section("retries")
    parser.add_section("retries")
    for port, count in sorted(retries.items(), key=lambda item: item[0].lower()):
        name = (port or "").strip()
        if not name:
            continue
        parser.set("retries", name, str(clamp_recv_retries(count)))
    save_config(parser)
    return parser


def _unique_port_names(ports: list[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for port in ports:
        name = (port or "").strip()
        if not name or name == UNUSED_PORT or name in seen:
            continue
        seen.add(name)
        ordered.append(name)
    return ordered


def parse_baud(text: str, fallback: int = 9600) -> int:
    try:
        value = int(float(str(text).strip()))
    except (TypeError, ValueError):
        return fallback
    return value if value > 0 else fallback


def load_port_baudrates(parser: configparser.ConfigParser | None = None) -> dict[str, int]:
    parser = parser or load_config()
    fallback = get_int(parser, "serial", "baudrate", 9600)
    mapping: dict[str, int] = {}
    if not parser.has_section("baud"):
        return mapping
    for key, value in parser.items("baud"):
        port = key.strip()
        if not port:
            continue
        mapping[port] = parse_baud(value, fallback)
    return mapping


def baud_for_port(port: str, baudrates: dict[str, int], fallback: int) -> int:
    if port in baudrates:
        return baudrates[port]
    lowered = {name.lower(): rate for name, rate in baudrates.items()}
    return lowered.get(port.lower(), fallback)


def load_port_encodings(parser: configparser.ConfigParser | None = None) -> dict[str, str]:
    parser = parser or load_config()
    mapping: dict[str, str] = {}
    if not parser.has_section("encoding"):
        return mapping
    for key, value in parser.items("encoding"):
        port = key.strip()
        if not port:
            continue
        mapping[port] = encoding_label(value)
    return mapping


def encoding_for_port(
    port: str,
    encodings: dict[str, str],
    fallback: str = DEFAULT_ENCODING,
) -> str:
    if port in encodings:
        return encoding_label(encodings[port])
    lowered = {name.lower(): label for name, label in encodings.items()}
    return encoding_label(lowered.get(port.lower(), fallback))


def load_port_timeouts(parser: configparser.ConfigParser | None = None) -> dict[str, float]:
    parser = parser or load_config()
    mapping: dict[str, float] = {}
    if not parser.has_section("timeout"):
        return mapping
    for key, value in parser.items("timeout"):
        port = key.strip()
        if not port:
            continue
        mapping[port] = clamp_recv_wait(value)
    return mapping


def timeout_for_port(port: str, timeouts: dict[str, float], fallback: float) -> float:
    if port in timeouts:
        return clamp_recv_wait(timeouts[port])
    lowered = {name.lower(): wait for name, wait in timeouts.items()}
    if port.lower() in lowered:
        return clamp_recv_wait(lowered[port.lower()])
    return clamp_recv_wait(fallback)


def load_port_retries(parser: configparser.ConfigParser | None = None) -> dict[str, int]:
    parser = parser or load_config()
    mapping: dict[str, int] = {}
    if not parser.has_section("retries"):
        return mapping
    for key, value in parser.items("retries"):
        port = key.strip()
        if not port:
            continue
        mapping[port] = clamp_recv_retries(value)
    return mapping


def retries_for_port(port: str, retries: dict[str, int], fallback: int) -> int:
    if port in retries:
        return clamp_recv_retries(retries[port])
    lowered = {name.lower(): count for name, count in retries.items()}
    if port.lower() in lowered:
        return clamp_recv_retries(lowered[port.lower()])
    return clamp_recv_retries(fallback)


def load_raw_length(parser: configparser.ConfigParser | None = None) -> int:
    parser = parser or load_config()
    return clamp_raw_length(get_int(parser, "channels", "raw_length", DEFAULT_RAW_LENGTH))


def load_lamp_ports(parser: configparser.ConfigParser | None = None) -> list[str]:
    parser = parser or load_config()
    if parser.has_section("com"):
        return _unique_port_names(get(parser, "com", "lamp", "").split(","))
    raw = _port_ini_lookup("lamp")
    if raw:
        return _unique_port_names(raw.split(","))
    return []


def load_com_enabled(parser: configparser.ConfigParser | None = None) -> bool:
    parser = parser or load_config()
    if parser.has_option("com", "use"):
        raw = get(parser, "com", "use", "1")
    else:
        looked = _port_ini_lookup("use")
        raw = looked if looked is not None else "1"
    if raw == "":
        return True
    return raw.strip().lower() not in ("0", "false", "off", "no")


def _port_ini_lookup(name: str) -> str | None:
    path = PORT_INI_PATH if PORT_INI_PATH.exists() else None
    if path is None and LEGACY_COM_MAP_PATH.exists():
        path = LEGACY_COM_MAP_PATH
    if path is None:
        return None
    try:
        text = path.read_text(encoding="utf-8")
    except OSError:
        return None
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or line.startswith(";"):
            continue
        if "=" not in line:
            continue
        key, value = line.split("=", 1)
        if key.strip().lower() == name:
            return value.strip()
    return None


def _sequence_from_done(value: Any) -> list[Any] | None:
    if isinstance(value, (list, tuple)):
        return list(value)
    if isinstance(value, (str, bytes, bytearray, dict)):
        return None
    to_list = getattr(value, "tolist", None)
    if callable(to_list):
        try:
            converted = to_list()
        except Exception:
            converted = None
        if isinstance(converted, list):
            return converted
        if isinstance(converted, tuple):
            return list(converted)
    if hasattr(value, "__iter__"):
        try:
            return list(value)
        except TypeError:
            return None
    return None


JUDGE_KEY = "判定"
JUDGE_ALIASES = {"判定", "ready"}


def judge_label(ok: Any) -> str:
    if isinstance(ok, str):
        text = ok.strip().upper()
        if text in ("TRUE", "FALSE"):
            return text
        if text in ("1", "YES", "OK"):
            return "TRUE"
        if text in ("0", "NO", ""):
            return "FALSE"
    if ok is True or ok == 1:
        return "TRUE"
    if ok is False or ok == 0:
        return "FALSE"
    return "TRUE" if bool(ok) else "FALSE"


def with_judge_last(values: dict[str, Any], fallback: Any | None = None) -> dict[str, Any]:
    found: Any = None
    out: dict[str, Any] = {}
    for key, value in values.items():
        name = str(key)
        if name in JUDGE_ALIASES:
            found = value
            continue
        out[name] = value
    flag = found if found is not None else fallback
    if flag is not None:
        out[JUDGE_KEY] = judge_label(flag)
    return out


def done_log_values(value: Any) -> dict[str, Any] | list[Any]:
    if value is None:
        return {"event": "done"}
    if isinstance(value, dict):
        return {str(key): item for key, item in value.items()}
    items = _sequence_from_done(value)
    if items is not None:
        return items
    return {"done": value}


def serial_options(parser: configparser.ConfigParser) -> dict[str, Any]:
    return {
        "baudrate": get_int(parser, "serial", "baudrate", 9600),
        "timeout": clamp_recv_wait(get(parser, "serial", "timeout", str(DEFAULT_RECV_WAIT))),
        "retries": clamp_recv_retries(get(parser, "serial", "retries", str(DEFAULT_RECV_RETRIES))),
        "bytesize": get_int(parser, "serial", "bytesize", 8),
        "parity": get(parser, "serial", "parity", "N") or "N",
        "stopbits": float(get(parser, "serial", "stopbits", "1") or 1),
    }


@dataclass
class CuiEntry:
    ui_name: str
    args: list[list[str]] = field(default_factory=list)
    title: str = ""
    texts: list[str] = field(default_factory=list)
    variables: dict[str, str] = field(default_factory=dict)
    headings: list[str] = field(default_factory=list)


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".")
    return str(value).strip()


def _is_empty_row(row: list[str]) -> bool:
    return not any(cell.strip() for cell in row)


def _is_cui_comment_mark(text: str) -> bool:
    first = str(text or "").strip()
    return first.startswith("#") or first.startswith("＃") or first.startswith("♯")


def _is_cui_comment_row(row: list[str]) -> bool:
    if not row:
        return False
    return _is_cui_comment_mark(row[0])


def _is_cui_header_row(arg_row: list[str]) -> bool:
    """変数・ラベルの見出し行。先頭の非空セルが # / ＃ / ♯ で始まる。"""
    for cell in arg_row:
        text = str(cell).strip()
        if not text:
            continue
        return _is_cui_comment_mark(text)
    return False


def _heading_from_row(row: list[str]) -> list[str]:
    cells: list[str] = []
    seen_mark = False
    for cell in row:
        text = str(cell).strip()
        if not seen_mark:
            if not text:
                continue
            if not _is_cui_comment_mark(text):
                return []
            seen_mark = True
            leftover = text.lstrip("#＃♯").strip()
            if leftover:
                cells.append(leftover)
            continue
        cells.append(text)
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _trim_grid(rows: list[list[str]]) -> list[list[str]]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    used_cols = [index for index in range(width) if any(padded[r][index].strip() for r in range(len(padded)))]
    if not used_cols:
        return []
    trimmed = [[row[index] for index in used_cols] for row in padded]
    body = [row for row in trimmed if any(cell.strip() for cell in row)]
    return strip_trailing_cui_empties(body)


def _parse_variables(args: list[list[str]]) -> dict[str, str]:
    variables: dict[str, str] = {}
    for row in args:
        for cell in row:
            if "=" not in cell:
                continue
            key, value = cell.split("=", 1)
            key = key.strip()
            if key:
                variables[key] = value.strip()
    return variables


def _derive_texts(args: list[list[str]]) -> list[str]:
    texts: list[str] = []
    for row in args:
        for cell in row:
            if not str(cell).strip():
                continue
            if "=" in cell:
                continue
            texts.append(cell)
    return texts


def _read_csv_grid(path: Path) -> list[list[str]]:
    data = path.read_bytes()
    if data.startswith(b"\xef\xbb\xbf"):
        text = data.decode("utf-8-sig")
    else:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("cp932")
    rows: list[list[str]] = []
    with io.StringIO(text, newline="") as handle:
        for row in csv.reader(handle):
            rows.append([_cell_str(cell) for cell in row])
    return rows


def _read_excel_grid(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append([_cell_str(cell) for cell in row])
        return rows
    finally:
        workbook.close()


def read_grid(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return _read_excel_grid(path)
    return _read_csv_grid(path)


def path_for_config(path: Path) -> str:
    resolved = path.expanduser()
    if not resolved.is_absolute():
        resolved = ROOT / resolved
    resolved = resolved.resolve()
    try:
        return str(resolved.relative_to(ROOT.resolve())).replace("\\", "/")
    except ValueError:
        return str(resolved)


def resolve_user_path(raw: str, fallback: str) -> Path:
    text = (raw or "").strip() or fallback
    path = Path(text).expanduser()
    if not path.is_absolute():
        path = ROOT / path
    return path


def resolve_log_directory(parser: configparser.ConfigParser | None = None) -> Path:
    source = parser or load_config()
    return resolve_user_path(get(source, "log", "directory", "logs"), "logs")


def resolve_backup_log_directory(parser: configparser.ConfigParser | None = None) -> Path | None:
    source = parser or load_config()
    raw = get(source, "log", "backup", "")
    if not raw:
        return None
    return resolve_user_path(raw, raw)


def resolve_cui_path(path: Path | None = None) -> Path:
    if path is not None:
        return path
    configured = get(load_config(), "cui", "file", "")
    if configured:
        candidate = resolve_user_path(configured, configured)
        if candidate.exists():
            return candidate
        raise FileNotFoundError(f"CUIファイルが見つかりません: {candidate}")
    for name in CUI_CANDIDATES:
        candidate = ROOT / name
        if candidate.exists():
            return candidate
    raise FileNotFoundError("CUIファイルが見つかりません（cui.xlsx / cui.csv）")


def parse_grid(rows: list[list[str]]) -> list[CuiEntry]:
    entries: list[CuiEntry] = []
    current: CuiEntry | None = None
    current_rows: list[list[str]] = []

    def flush() -> None:
        nonlocal current, current_rows
        if current is None:
            return
        args = _trim_grid(current_rows)
        current.args = args
        current.texts = _derive_texts(args)
        current.variables = _parse_variables(args)
        entries.append(current)
        current = None
        current_rows = []

    for row in rows:
        normalized = [_cell_str(cell) for cell in row]
        if _is_empty_row(normalized):
            continue
        if current is not None and not current.headings:
            if _is_cui_comment_row(normalized) and str(normalized[0]).strip().startswith("♯"):
                current.headings = _heading_from_row(normalized)
                continue
            arg_head = normalized[1:] if not str(normalized[0]).strip() else []
            if arg_head and _is_cui_header_row(arg_head):
                current.headings = _heading_from_row(arg_head)
                continue
        if _is_cui_comment_row(normalized):
            continue
        name = normalized[0].strip()
        if name:
            flush()
            current = CuiEntry(ui_name=name)
            current.title = normalized[1].strip() if len(normalized) > 1 else ""
            current_rows = []
        elif current is not None:
            arg_row = normalized[1:]
            if _is_cui_header_row(arg_row):
                if not current.headings:
                    current.headings = _heading_from_row(arg_row)
                continue
            current_rows.append(arg_row)
    flush()
    return entries


def load_cui(path: Path | None = None) -> list[CuiEntry]:
    return parse_grid(read_grid(resolve_cui_path(path)))


def _is_judge_row(row: list[str]) -> bool:
    return bool(row) and str(row[0]).strip() == JUDGE_KEY


def _judge_row_value(row: list[str]) -> str:
    if len(row) < 2:
        return ""
    return str(row[1]).strip().upper()


def _record_true_stamp(stamp: str) -> str:
    text = str(stamp or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y/%m/%d %H:%M:%S")
        except ValueError:
            continue
    return text


def _last_true_unit(rows: list[tuple[str, list[str]]]) -> list[tuple[str, list[str]]]:
    last: list[tuple[str, list[str]]] = []
    current: list[tuple[str, list[str]]] = []
    for row in rows:
        current.append(row)
        if not _is_judge_row(row[1]):
            continue
        if _judge_row_value(row[1]) == "TRUE":
            last = current
        current = []
    return last


def _true_as_stamp(rows: list[tuple[str, list[str]]]) -> list[list[str]]:
    out: list[list[str]] = []
    for stamp, data in rows:
        cells = list(data)
        if _is_judge_row(cells) and _judge_row_value(cells) == "TRUE":
            shown = _record_true_stamp(stamp)
            if shown:
                if len(cells) >= 2:
                    cells[1] = shown
                else:
                    cells.append(shown)
        out.append(cells)
    return out


def _log_blocks(records: list[list[str]]) -> list[tuple[str, str, list[tuple[str, list[str]]]]]:
    blocks: list[tuple[str, str, list[tuple[str, list[str]]]]] = []
    for row in records:
        stamp = str(row[0]) if row else ""
        name = str(row[1]) if len(row) > 1 else ""
        title = str(row[2]) if len(row) > 2 else ""
        data = [str(cell) for cell in row[3:]]
        while data and not str(data[-1]).strip():
            data.pop()
        if not name and not data:
            continue
        item = (stamp, data)
        if not blocks or blocks[-1][0] != name or blocks[-1][1] != title:
            blocks.append((name, title, [item]))
        else:
            blocks[-1][2].append(item)
    return blocks


def _assign_log_blocks(entries: list[CuiEntry], records: list[list[str]]) -> list[list[list[str]]]:
    unused = _log_blocks(records)
    assigned: list[list[list[str]]] = []
    for entry in entries:
        found: int | None = None
        title = str(entry.title or "")
        for index, (name, heading, _rows) in enumerate(unused):
            if name == entry.ui_name and heading == title:
                found = index
                break
        if found is None:
            for index, (name, _heading, _rows) in enumerate(unused):
                if name == entry.ui_name:
                    found = index
                    break
        if found is None:
            assigned.append([])
        else:
            assigned.append(_true_as_stamp(_last_true_unit(unused.pop(found)[2])))
    return assigned


def _record_save_name(cui_name: str, *, manual: bool = False) -> str:
    name = Path(cui_name or "cui").stem or "cui"
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    suffix = "_manual" if manual else ""
    return f"{name}_{stamp}{suffix}.csv"


def _record_log_cells(row: list[str]) -> list[str]:
    if is_heading_row(row):
        return heading_cells(row)
    return [str(cell) for cell in row]


def _record_view_layout(
    entries: list[CuiEntry],
    assigned: list[list[list[str]]],
) -> tuple[list[list[str]], set[tuple[int, int]], set[int], set[int]]:
    """記録確認と同じ格子。戻り値は行、見出し座標、隙間列、空行。"""
    grid: list[list[str]] = []
    headers: set[tuple[int, int]] = set()
    gaps: set[int] = set()
    spacers: set[int] = set()

    def put_cell(row: int, col: int, text: str, header: bool = False) -> None:
        while len(grid) <= row:
            grid.append([])
        line = grid[row]
        while len(line) <= col:
            line.append("")
        line[col] = text
        if header:
            headers.add((row, col))

    if not entries:
        put_cell(0, 0, "CUI に UI がありません。")
        return grid, headers, gaps, spacers

    for index, (entry, log_rows) in enumerate(zip(entries, assigned)):
        start = len(grid)
        put_cell(start, 0, str(entry.title or ""), header=True)
        put_cell(start, 1, entry.ui_name, header=True)
        data_row = start + 1
        shown = [_record_log_cells(row) for row in log_rows]
        log_width = max((len(row) for row in shown), default=1) or 1
        gap = log_width
        gaps.add(gap)
        set_col = gap + 1
        for offset, row in enumerate(shown):
            heading = is_heading_row(log_rows[offset])
            for col, cell in enumerate(row):
                put_cell(data_row + offset, col, cell, header=heading)
        set_rows = 0
        if entry.headings:
            for col, cell in enumerate(entry.headings):
                put_cell(data_row, set_col + col, str(cell), header=True)
            set_rows = 1
        for offset, row in enumerate(entry.args):
            for col, cell in enumerate(row):
                put_cell(data_row + set_rows + offset, set_col + col, str(cell))
        used = max(len(shown), set_rows + len(entry.args), 1)
        while len(grid) < data_row + used:
            grid.append([])
        if index < len(entries) - 1:
            spacers.add(len(grid))
            grid.append([])
    return grid, headers, gaps, spacers


def _write_record_view(path: Path, grid: list[list[str]]) -> None:
    width = max((len(row) for row in grid), default=0)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        for row in grid:
            if not row:
                writer.writerow([])
                continue
            padded = list(row) + [""] * (width - len(row))
            writer.writerow(padded)


class MainApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.config = load_config()
        self.entries: list[CuiEntry] = []
        self.index = 0
        self.current_ui: BaseUI | None = None
        self.current_entry: CuiEntry | None = None
        self._last_logged: dict[str, Any] | None = None

        log_dir = resolve_log_directory(self.config)
        self.logger = self._make_logger(log_dir, "")
        self.channels = ChannelHub(on_data=self._on_serial_data)
        self.design = load_design()
        self._ensure_com_config()

        self._cui_load_win: tk.Toplevel | None = None
        self._startup_cui_started = False
        self._build_window()
        self._show_current_ui()
        self._tick_condition()
        self.root.after(1, self._after_main_shown)

    def _ensure_com_config(self) -> None:
        if self.config.has_section("com") and _channel_keys_present(self.config):
            return
        self.config = save_channel_ports(
            load_channel_ports(self.config),
            channel_max(self.config),
            load_lamp_ports(self.config),
            load_com_enabled(self.config),
            parser=self.config,
        )

    def _build_window(self) -> None:
        title = get(self.config, "window", "title", "kamishivy")
        width = get_int(self.config, "window", "width", 960)
        height = get_int(self.config, "window", "height", 640)
        self.root.title(title)
        self.root.geometry(f"{width}x{height}")
        self.root.minsize(640, 480)
        self.root.protocol("WM_DELETE_WINDOW", self._on_exit)

        style = ttk.Style(self.root)
        if "clam" in style.theme_names():
            style.theme_use("clam")
        elif "vista" in style.theme_names():
            style.theme_use("vista")
        self._style = style
        style.configure("Header.TFrame", padding=8)
        style.configure("Footer.TFrame", padding=10)

        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(1, weight=1)

        header = ttk.Frame(self.root, style="Header.TFrame")
        header.grid(row=0, column=0, sticky="ew")
        header.columnconfigure(1, weight=1)
        self.header = header

        left = ttk.Frame(header, style="HeaderInner.TFrame")
        left.grid(row=0, column=0, sticky="w")
        bg = style.lookup("TFrame", "background") or self.root.cget("background")
        self._lamp_bg = bg
        self.connect_button = tk.Button(left, text="接続", command=self._toggle_connection)
        self.connect_button.pack(side="left")
        self.com_lamp_frame = ttk.Frame(left, style="HeaderInner.TFrame")
        self.com_lamp_frame.pack(side="left", padx=(16, 0))
        self._com_lamps: dict[str, tuple[tk.Canvas, int]] = {}
        self._overall_lamp = LAMP_OFF
        self.header_left = left
        self._apply_com_usage()

        right = ttk.Frame(header, style="HeaderInner.TFrame")
        right.grid(row=0, column=2, sticky="e")
        self.header_right = right
        self.page_label = ttk.Label(right, text="", style="Header.TLabel")
        menu_button = tk.Menubutton(right, text="メニュー", direction="below")
        menu_button.pack(side="right")
        menu = tk.Menu(menu_button, tearoff=False)
        menu.add_command(label="通信", command=self._open_com_dialog)
        menu.add_command(label="通信データの加工", command=self._open_parse_dialog)
        menu.add_command(label="表示のカスタマイズ", command=self._open_display_dialog)
        menu.add_command(label="設定", command=self._open_settings_dialog)
        menu.add_command(label="次へ", command=self._go_next_from_menu)
        menu.add_command(label="記録確認", command=self._open_record_dialog)
        menu.add_command(label="戻る", command=self._go_back)
        menu.add_command(label="UI一覧", command=self._open_ui_list_dialog)
        menu.add_separator()
        menu.add_command(label="終了", command=self._on_exit)
        menu_button.configure(menu=menu)
        self.menu = menu
        self.menu_button = menu_button

        self.content = ttk.Frame(self.root, padding=12)
        self.content.grid(row=1, column=0, sticky="nsew")
        self.content.columnconfigure(0, weight=1)
        self.content.rowconfigure(0, weight=1)

        self.footer = ttk.Frame(self.root, style="Footer.TFrame")
        self.footer.grid(row=2, column=0, sticky="ew")
        self.next_holder = tk.Frame(self.footer, width=160, height=44)
        self.next_holder.pack_propagate(False)
        self.next_button = tk.Button(self.next_holder, text="次へ", command=self._go_next)
        self.next_button.pack(fill="both", expand=True)
        self._set_next_enabled(False)
        self._apply_design()
        self._apply_window_mode()
        self._layout_next_button()
        self.root.bind("<Configure>", self._on_root_configure)

    def _on_root_configure(self, event: tk.Event) -> None:
        if event.widget is not self.root:
            return
        self._layout_next_button()

    def _tone_hex(self, color: str, delta: int) -> str:
        rgb = hex_to_rgb(color, (225, 225, 225))
        return rgb_hex(tuple(max(0, min(255, channel + delta)) for channel in rgb))

    def _apply_tk_button(
        self,
        widget: tk.Button | tk.Menubutton,
        *,
        font: tuple,
        bg: str,
        fg: str,
        border: str,
        border_w: int,
        enabled_border: bool,
        disabled_fg: str | None = None,
        padx: int = 8,
        pady: int = 2,
    ) -> None:
        active = self._tone_hex(bg, -25)
        widget.configure(
            font=font,
            bg=bg,
            fg=fg,
            activebackground=active,
            activeforeground=fg,
            disabledforeground=disabled_fg or fg,
            highlightbackground=border,
            highlightcolor=border,
            highlightthickness=border_w,
            bd=border_w,
            relief="solid" if enabled_border else "raised",
            padx=padx,
            pady=pady,
        )

    def _apply_design(self) -> None:
        d = self.design
        bg = str(d.get("background") or "#f0f0f0")
        fg = str(d.get("text") or "#1e1e1e")
        button = str(d.get("button") or "#e1e1e1")
        gadget = str(d.get("gadget") or "#ffffff")
        textbox = str(d.get("textbox") or "#ffffff")
        border = str(d.get("border") or "#555555")
        menu_bg = str(d.get("menu") or button)
        family = str(d.get("font_family") or DEFAULTS["font_family"])
        size = int(d.get("font_size") or DEFAULTS["font_size"])
        font = font_tuple(family, size, d.get("font_bold"))
        header_bg = str(d.get("header_background") or bg)
        header_fg = str(d.get("header_text") or fg)
        header_button = str(d.get("header_button") or button)
        header_family = str(d.get("header_font_family") or family)
        header_size = int(d.get("header_font_size") or size)
        header_font = font_tuple(header_family, header_size, d.get("header_font_bold"))
        menu_fg = str(d.get("menu_text") or header_fg)
        footer_bg = str(d.get("footer_background") or bg)
        footer_fg = str(d.get("footer_text") or fg)
        footer_button = str(d.get("footer_button") or button)
        footer_family = str(d.get("footer_font_family") or family)
        footer_size = int(d.get("footer_font_size") or size)
        button_fg = str(d.get("button_text") or fg)
        button_font = d.get("button_font") or font_tuple(
            family, int(d.get("button_font_size") or size), d.get("font_bold")
        )
        header_button_fg = str(d.get("header_button_text") or header_fg)
        header_button_font = d.get("header_button_font") or font_tuple(
            header_family, int(d.get("header_button_font_size") or header_size), d.get("header_font_bold")
        )
        footer_button_fg = str(d.get("footer_button_text") or footer_fg)
        footer_button_font = d.get("footer_button_font") or font_tuple(
            footer_family, int(d.get("footer_button_font_size") or footer_size), d.get("footer_font_bold")
        )
        enabled_border = int(d.get("border_enabled") or 0) != 0
        border_w = int(d.get("border_width") or 0) if enabled_border else 0
        self.root.configure(bg=bg)
        self._lamp_bg = header_bg
        try:
            margin = int(d.get("content_margin", DEFAULTS["content_margin"]))
        except (TypeError, ValueError):
            margin = int(DEFAULTS["content_margin"])
        margin = max(0, min(200, margin))
        try:
            self.content.configure(padding=(margin, 12, margin, 12))
        except tk.TclError:
            pass
        style = getattr(self, "_style", None) or ttk.Style(self.root)
        style.configure(".", font=font, background=bg, foreground=fg)
        style.configure("TFrame", background=bg)
        style.configure("Header.TFrame", background=header_bg, padding=8)
        style.configure("HeaderInner.TFrame", background=header_bg, padding=0)
        style.configure("Header.TLabel", background=header_bg, foreground=header_fg, font=header_font)
        style.configure("Footer.TFrame", background=footer_bg, padding=10)
        style.configure("TLabel", background=bg, foreground=fg, font=font, borderwidth=0, relief="flat")
        style.configure("TButton", background=button, foreground=button_fg, font=button_font)
        style.map("TButton", background=[("active", button), ("disabled", gadget)])
        indicator = max(14, int(round(size * 1.4)))
        style.configure("TCheckbutton", background=bg, foreground=fg, font=font, indicatorsize=indicator)
        style.configure("TRadiobutton", background=bg, foreground=fg, font=font, indicatorsize=indicator)
        dialog_indicator = max(10, int(round(self._dialog_font_size() * 1.1)))
        style.configure("Dialog.TCheckbutton", indicatorsize=dialog_indicator)
        style.configure("Dialog.TRadiobutton", indicatorsize=dialog_indicator)
        style.configure("TMenubutton", background=header_button, foreground=header_button_fg, font=header_button_font)
        style.map("TMenubutton", background=[("active", header_button), ("pressed", header_button)])
        style.configure("TLabelframe", background=bg, foreground=fg)
        style.configure("TLabelframe.Label", background=bg, foreground=fg, font=font)
        style.configure("Settings.TLabelframe", background=bg, foreground=fg)
        style.configure("Settings.TLabelframe.Label", background=bg, foreground=fg, font=font)
        style.configure("TNotebook", background=bg)
        style.configure("TNotebook.Tab", background=button, foreground=fg, font=font)
        style.configure("TEntry", fieldbackground=textbox, foreground=fg, font=font)
        style.configure("TCombobox", fieldbackground=textbox, background=gadget, foreground=fg, font=font)
        fill_bg = [("active", bg), ("selected", bg), ("pressed", bg), ("disabled", bg)]
        style.map("TLabel", background=fill_bg)
        style.map("Header.TLabel", background=[("active", header_bg), ("disabled", header_bg)])
        style.map("TCheckbutton", background=fill_bg)
        style.map("TRadiobutton", background=fill_bg)
        style.map("TLabelframe.Label", background=fill_bg)
        style.map("Settings.TLabelframe.Label", background=fill_bg)
        style.map("TFrame", background=fill_bg)
        if border_w:
            style.configure("TLabelframe", borderwidth=border_w, relief="solid")
        else:
            style.configure("TLabelframe", borderwidth=1, relief="groove")
        for widget, style_name in (
            (self.header, "Header.TFrame"),
            (self.header_left, "HeaderInner.TFrame"),
            (self.header_right, "HeaderInner.TFrame"),
            (self.com_lamp_frame, "HeaderInner.TFrame"),
            (self.content, "TFrame"),
            (self.footer, "Footer.TFrame"),
        ):
            try:
                widget.configure(style=style_name)
            except tk.TclError:
                pass
        self._apply_tk_button(
            self.menu_button,
            font=header_button_font,
            bg=header_button,
            fg=header_button_fg,
            border=border,
            border_w=border_w,
            enabled_border=enabled_border,
        )
        self._apply_tk_button(
            self.connect_button,
            font=header_button_font,
            bg=header_button,
            fg=header_button_fg,
            border=border,
            border_w=border_w,
            enabled_border=enabled_border,
        )
        self.menu.configure(
            font=header_font,
            bg=menu_bg,
            fg=menu_fg,
            activebackground=self._tone_hex(menu_bg, -25),
            activeforeground=menu_fg,
            relief="solid" if enabled_border else "raised",
            bd=max(1, border_w),
        )
        self.next_holder.configure(bg=footer_bg)
        self.footer.configure(style="Footer.TFrame")
        self._apply_tk_button(
            self.next_button,
            font=footer_button_font,
            bg=footer_button,
            fg=footer_button_fg,
            border=border,
            border_w=border_w,
            enabled_border=enabled_border,
            disabled_fg=str(d.get("footer_text_disabled") or footer_button_fg),
            padx=0,
            pady=0,
        )
        self.next_button.configure(relief="solid" if enabled_border else "flat")
        self._apply_next_button_colors()
        if load_com_enabled(self.config):
            self._rebuild_com_lamps()

    def _apply_window_mode(self) -> None:
        fullscreen = get_int(self.config, "display", "fullscreen", 0) != 0
        hide_close = get_int(self.config, "display", "hide_close", 0) != 0
        self.root.attributes("-fullscreen", False)
        try:
            self.root.state("normal")
        except tk.TclError:
            pass
        if fullscreen and hide_close:
            self.root.attributes("-fullscreen", True)
            return
        if fullscreen:
            try:
                self.root.state("zoomed")
            except tk.TclError:
                self.root.attributes("-fullscreen", True)
            return
        width = get_int(self.config, "window", "width", 960)
        height = get_int(self.config, "window", "height", 640)
        self.root.geometry(f"{width}x{height}")

    def _next_pixel_size(self) -> tuple[int, int]:
        mode = (get(self.config, "display", "next_size_mode", NEXT_SIZE_MODE_DEFAULT) or NEXT_SIZE_MODE_DEFAULT).lower()
        width_value = get_float(self.config, "display", "next_width", NEXT_WIDTH_DEFAULT)
        height_value = get_float(self.config, "display", "next_height", NEXT_HEIGHT_DEFAULT)
        win_w = max(1, int(self.root.winfo_width() or get_int(self.config, "window", "width", 960)))
        win_h = max(1, int(self.root.winfo_height() or get_int(self.config, "window", "height", 640)))
        if mode == "px":
            width = int(width_value)
            height = int(height_value)
            return (max(24, min(win_w - 24, width)), max(20, min(win_h - 24, height)))
        width = int(win_w * max(0, width_value) / 100)
        height = int(win_h * max(0, height_value) / 100)
        return (max(8, min(win_w - 24, width)), max(8, min(win_h - 24, height)))

    def _layout_next_button(self) -> None:
        if not getattr(self, "next_holder", None):
            return
        width, height = self._next_pixel_size()
        if getattr(self, "_next_size", None) != (width, height):
            self._next_size = (width, height)
            self.next_holder.configure(width=width, height=height)
        align = (get(self.config, "display", "next_align", "center") or "center").lower()
        if getattr(self, "_next_align", None) == align and self.next_holder.winfo_manager():
            return
        self._next_align = align
        self.next_holder.pack_forget()
        padx, pady = 12, 8
        if align == "left":
            self.next_holder.pack(side="left", padx=padx, pady=pady)
        elif align == "right":
            self.next_holder.pack(side="right", padx=padx, pady=pady)
        else:
            self.next_holder.pack(anchor="center", padx=padx, pady=pady)

    def _apply_com_usage(self) -> None:
        enabled = load_com_enabled(self.config)
        if not enabled and self.channels.connected:
            self.channels.disconnect()
            self.connect_button.config(text="接続")
            self._set_lamp(LAMP_OFF)
            self._refresh_condition()
        self.connect_button.pack_forget()
        self.com_lamp_frame.pack_forget()
        if not enabled:
            for child in self.com_lamp_frame.winfo_children():
                child.destroy()
            self._com_lamps = {}
            return
        self.connect_button.pack(side="left")
        self.com_lamp_frame.pack(side="left", padx=(16, 0))
        self._rebuild_com_lamps()

    def _set_lamp(self, color: str) -> None:
        self._overall_lamp = color
        self._refresh_com_lamps()

    def _rebuild_com_lamps(self) -> None:
        for child in self.com_lamp_frame.winfo_children():
            child.destroy()
        self._com_lamps = {}
        bg = getattr(self, "_lamp_bg", None) or self.root.cget("background")
        for port in load_lamp_ports(self.config):
            cell = ttk.Frame(self.com_lamp_frame, style="HeaderInner.TFrame")
            cell.pack(side="left", padx=(0, 12))
            canvas = tk.Canvas(cell, width=18, height=18, highlightthickness=0, bd=0, background=bg)
            lamp_id = canvas.create_oval(2, 2, 16, 16, fill=LAMP_OFF, outline="#555555")
            canvas.pack(side="left")
            ttk.Label(cell, text=port, style="Header.TLabel").pack(side="left", padx=(4, 0))
            self._com_lamps[port] = (canvas, lamp_id)
        self._refresh_com_lamps()

    def _com_lamp_color(self, port: str) -> str:
        overall = getattr(self, "_overall_lamp", LAMP_OFF)
        if overall == LAMP_WAIT:
            return LAMP_WAIT
        health = self.channels.port_health()
        if port in health:
            return LAMP_ON if health[port] == "on" else LAMP_ERROR
        mapping = load_channel_ports(self.config)
        failed = {
            mapping[number]
            for number in self.channels.last_errors
            if mapping.get(number)
        }
        if port in failed:
            return LAMP_ERROR
        return LAMP_OFF

    def _refresh_com_lamps(self) -> None:
        lamps = getattr(self, "_com_lamps", None)
        if not lamps:
            return
        for port, (canvas, lamp_id) in lamps.items():
            canvas.itemconfig(lamp_id, fill=self._com_lamp_color(port))

    def _dialog_font_size(self) -> int:
        try:
            menu_size = int(self.design.get("header_font_size") or DEFAULTS["header_font_size"])
        except (TypeError, ValueError):
            menu_size = int(DEFAULTS["header_font_size"])
        return max(1, int(round(menu_size * 0.95)))

    def _dialog_look(self) -> dict[str, Any]:
        base = default_design()
        family = str(base.get("font_family") or DEFAULTS["font_family"])
        size = self._dialog_font_size()
        return {
            "font": font_tuple(family, size, 0),
            "size": size,
            "bg": str(base.get("background") or rgb_hex(DEFAULTS["background"])),
            "fg": str(base.get("text") or rgb_hex(DEFAULTS["text"])),
            "button": str(base.get("button") or rgb_hex(DEFAULTS["button"])),
            "button_fg": str(base.get("button_text") or rgb_hex(DEFAULTS["button_text"])),
            "gadget": str(base.get("gadget") or rgb_hex(DEFAULTS["gadget"])),
            "textbox": str(base.get("textbox") or rgb_hex(DEFAULTS["textbox"])),
        }

    def _ensure_dialog_styles(self) -> None:
        look = self._dialog_look()
        style = getattr(self, "_style", None) or ttk.Style(self.root)
        font = look["font"]
        size = int(look["size"])
        bg = look["bg"]
        fg = look["fg"]
        button = look["button"]
        button_fg = look["button_fg"]
        gadget = look["gadget"]
        textbox = look["textbox"]
        indicator = max(10, int(round(size * 1.1)))
        row_h = max(22, int(round(size * 1.8)))
        style.configure("Dialog.TFrame", background=bg)
        style.configure("Dialog.TLabel", background=bg, foreground=fg, font=font, borderwidth=0, relief="flat")
        style.configure("Dialog.TButton", background=button, foreground=button_fg, font=font)
        style.map("Dialog.TButton", background=[("active", button), ("disabled", gadget)])
        style.configure("Dialog.TCheckbutton", background=bg, foreground=fg, font=font, indicatorsize=indicator)
        style.configure("Dialog.TRadiobutton", background=bg, foreground=fg, font=font, indicatorsize=indicator)
        style.configure("Dialog.TLabelframe", background=bg, foreground=fg, borderwidth=1, relief="groove")
        style.configure("Dialog.TLabelframe.Label", background=bg, foreground=fg, font=font)
        style.configure("Dialog.TNotebook", background=bg)
        style.configure("Dialog.TNotebook.Tab", background=button, foreground=fg, font=font)
        style.configure("Dialog.TEntry", fieldbackground=textbox, foreground=fg, font=font)
        style.configure("Dialog.TCombobox", fieldbackground=textbox, background=gadget, foreground=fg, font=font)
        style.configure("Dialog.TSpinbox", fieldbackground=textbox, foreground=fg, font=font)
        style.configure("Dialog.Treeview", background=textbox, foreground=fg, fieldbackground=textbox, font=font, rowheight=row_h)
        style.configure("Dialog.Treeview.Heading", background=button, foreground=fg, font=font)
        style.configure("Dialog.TScrollbar", background=button)
        fill_bg = [("active", bg), ("selected", bg), ("pressed", bg), ("disabled", bg)]
        style.map("Dialog.TLabel", background=fill_bg)
        style.map("Dialog.TCheckbutton", background=fill_bg)
        style.map("Dialog.TRadiobutton", background=fill_bg)
        style.map("Dialog.TLabelframe.Label", background=fill_bg)
        style.map("Dialog.TFrame", background=fill_bg)

    def _skin_dialog(self, win: tk.Toplevel) -> None:
        look = self._dialog_look()
        bg = look["bg"]
        font = look["font"]
        mapping = {
            "TFrame": "Dialog.TFrame",
            "TLabel": "Dialog.TLabel",
            "TButton": "Dialog.TButton",
            "TCheckbutton": "Dialog.TCheckbutton",
            "TRadiobutton": "Dialog.TRadiobutton",
            "TLabelframe": "Dialog.TLabelframe",
            "TNotebook": "Dialog.TNotebook",
            "TEntry": "Dialog.TEntry",
            "TCombobox": "Dialog.TCombobox",
            "TSpinbox": "Dialog.TSpinbox",
            "Treeview": "Dialog.Treeview",
            "TScrollbar": "Dialog.TScrollbar",
        }
        tk_fonts = {"Text", "Entry", "Button", "Label", "Checkbutton", "Radiobutton", "Listbox", "Spinbox"}
        try:
            win.configure(bg=bg)
        except tk.TclError:
            pass

        def walk(widget) -> None:
            try:
                cls = str(widget.winfo_class())
            except tk.TclError:
                return
            style_name = mapping.get(cls)
            if style_name:
                try:
                    widget.configure(style=style_name)
                except tk.TclError:
                    pass
            elif cls in tk_fonts:
                try:
                    widget.configure(font=font)
                except tk.TclError:
                    pass
                if cls == "Text":
                    try:
                        widget.configure(bg=look["textbox"], fg=look["fg"], insertbackground=look["fg"])
                    except tk.TclError:
                        pass
            for child in widget.winfo_children():
                walk(child)

        walk(win)

    def _use_default_dialog_look(self, win: tk.Toplevel) -> None:
        """メニューから開く窓は色などを表示のカスタマイズから取らない。文字サイズだけメニューバーの95%。"""
        self._ensure_dialog_styles()
        self._skin_dialog(win)

    def _open_display_dialog(self) -> None:
        win = tk.Toplevel(self.root)
        win.title("表示のカスタマイズ")
        win.transient(self.root)
        win.geometry("720x560")
        win.minsize(640, 420)
        win.grab_set()
        self._ensure_dialog_styles()

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)

        families = sorted(set(tkfont.families(self.root)), key=str.lower)
        font_vars = {
            "": tk.StringVar(value=str(self.design.get("font_family") or DEFAULTS["font_family"])),
            "header": tk.StringVar(
                value=str(self.design.get("header_font_family") or self.design.get("font_family") or DEFAULTS["font_family"])
            ),
            "footer": tk.StringVar(
                value=str(self.design.get("footer_font_family") or self.design.get("font_family") or DEFAULTS["font_family"])
            ),
        }
        size_vars = {
            "": tk.StringVar(value=str(int(self.design.get("font_size") or DEFAULTS["font_size"]))),
            "header": tk.StringVar(
                value=str(int(self.design.get("header_font_size") or self.design.get("font_size") or DEFAULTS["font_size"]))
            ),
            "footer": tk.StringVar(
                value=str(int(self.design.get("footer_font_size") or self.design.get("font_size") or DEFAULTS["footer_font_size"]))
            ),
        }
        bold_vars = {
            "": tk.BooleanVar(value=int(self.design.get("font_bold") or 0) != 0),
            "header": tk.BooleanVar(value=int(self.design.get("header_font_bold") or 0) != 0),
            "footer": tk.BooleanVar(value=int(self.design.get("footer_font_bold") or 0) != 0),
        }
        button_size_vars = {
            "": tk.StringVar(
                value=str(int(self.design.get("button_font_size") or self.design.get("font_size") or DEFAULTS["button_font_size"]))
            ),
            "header": tk.StringVar(
                value=str(
                    int(
                        self.design.get("header_button_font_size")
                        or self.design.get("header_font_size")
                        or DEFAULTS["header_button_font_size"]
                    )
                )
            ),
            "footer": tk.StringVar(
                value=str(
                    int(
                        self.design.get("footer_button_font_size")
                        or self.design.get("footer_font_size")
                        or DEFAULTS["footer_button_font_size"]
                    )
                )
            ),
        }
        title_font_var = tk.StringVar(
            value=str(self.design.get("title_font_family") or self.design.get("font_family") or DEFAULTS["title_font_family"])
        )
        title_size_var = tk.StringVar(
            value=str(int(self.design.get("title_font_size") or DEFAULTS["title_font_size"]))
        )
        margin_var = tk.StringVar(
            value=str(int(self.design.get("content_margin", DEFAULTS["content_margin"])))
        )
        color_labels = {
            "text": "文字色",
            "header_text": "文字色",
            "footer_text": "文字色",
            "footer_text_disabled": "押す前の文字色",
            "background": "背景色",
            "header_background": "背景色",
            "footer_background": "背景色",
            "button": "ボタンの色",
            "header_button": "ボタンの色",
            "footer_button": "ボタンの色",
            "button_text": "文字色",
            "header_button_text": "文字色",
            "footer_button_text": "文字色",
            "footer_button_disabled": "押す前の色",
            "gadget": "ガジェットの色",
            "textbox": "テキストボックスの背景色",
            "border": "縁取りの色",
            "menu": "メニューの色（ドロップダウン）",
            "menu_text": "メニューの文字色",
            "title_text": "文字色",
        }
        color_vars: dict[str, tk.StringVar] = {}
        swatches: dict[str, tk.Frame] = {}
        font_combos: dict[str, ttk.Combobox] = {}

        def color_value(key: str) -> str:
            fallback_key = key.split("_", 1)[-1] if "_" in key else key
            return str(self.design.get(key) or self.design.get(fallback_key) or rgb_hex((240, 240, 240)))

        def pick_color(key: str) -> None:
            current = color_vars[key].get()
            chosen = colorchooser.askcolor(color=current, parent=win, title=color_labels.get(key, key))
            if not chosen or not chosen[1]:
                return
            color_vars[key].set(chosen[1])
            swatches[key].configure(bg=chosen[1])

        def add_color_row(parent: ttk.Frame, key: str, label: str) -> None:
            hex_value = color_value(key)
            color_vars[key] = tk.StringVar(value=hex_value)
            block = ttk.Frame(parent)
            block.pack(fill="x", pady=(0, 8))
            ttk.Label(block, text=label).pack(anchor="w")
            line = ttk.Frame(block)
            line.pack(anchor="w", pady=(2, 0))
            swatch = tk.Frame(line, width=28, height=18, bg=hex_value, highlightthickness=1, highlightbackground="#888")
            swatch.pack(side="left")
            swatch.pack_propagate(False)
            swatches[key] = swatch
            ttk.Button(line, text="色を選ぶ", command=lambda k=key: pick_color(k)).pack(side="left", padx=(8, 0))

        def add_font_fields(parent: ttk.Frame, region: str) -> None:
            ttk.Label(parent, text="フォント").pack(anchor="w", pady=(0, 2))
            combo = ttk.Combobox(parent, textvariable=font_vars[region], values=families, width=22, state="readonly")
            combo.pack(anchor="w")
            font_combos[region] = combo
            if font_vars[region].get() not in families:
                combo.config(state="normal")
            ttk.Label(parent, text="文字の大きさ").pack(anchor="w", pady=(8, 2))
            ttk.Spinbox(parent, from_=8, to=72, textvariable=size_vars[region], width=8).pack(anchor="w")
            ttk.Checkbutton(parent, text="太字", variable=bold_vars[region]).pack(anchor="w", pady=(6, 0))
            add_color_row(parent, f"{region}_text" if region else "text", "文字色")

        def add_button_fields(parent: ttk.Frame, region: str, bg_key: str, text_key: str) -> None:
            add_color_row(parent, bg_key, "ボタンの色")
            add_color_row(parent, text_key, "文字色")
            ttk.Label(parent, text="文字の大きさ").pack(anchor="w", pady=(0, 2))
            ttk.Spinbox(parent, from_=8, to=72, textvariable=button_size_vars[region], width=8).pack(anchor="w")

        def section_box(parent: ttk.Widget, text: str, **pack) -> ttk.Frame:
            box = ttk.LabelFrame(parent, text=text, style="Dialog.TLabelframe")
            box.pack(**pack)
            inner = ttk.Frame(box)
            inner.pack(fill="both", expand=True, padx=12, pady=(8, 14))
            return inner

        def region_columns(parent: ttk.Frame) -> tuple[ttk.Frame, ttk.Frame, ttk.Frame]:
            row = ttk.Frame(parent)
            row.pack(fill="x", pady=(0, 8))
            inners: list[ttk.Frame] = []
            for index, title in enumerate(("メニューバー", "表示域", "フッター")):
                box = ttk.LabelFrame(row, text=title, style="Dialog.TLabelframe")
                box.pack(side="left", fill="both", expand=True, padx=(0 if index == 0 else 8, 0))
                inner = ttk.Frame(box)
                inner.pack(fill="both", expand=True, padx=12, pady=(8, 14))
                inners.append(inner)
            return inners[0], inners[1], inners[2]

        notebook = ttk.Notebook(frame)
        font_tab = ttk.Frame(notebook, padding=8)
        button_tab = ttk.Frame(notebook, padding=8)
        window_tab = ttk.Frame(notebook, padding=8)
        file_tab = ttk.Frame(notebook, padding=8)
        notebook.add(font_tab, text="フォント")
        notebook.add(button_tab, text="ボタン")
        notebook.add(window_tab, text="全体設定")
        notebook.add(file_tab, text="保存・読み出し")

        font_top, font_mid, font_bottom = region_columns(font_tab)
        add_font_fields(font_top, "header")
        add_font_fields(font_mid, "")
        ttk.Label(font_mid, text="左右のマージン").pack(anchor="w", pady=(8, 2))
        margin_row = ttk.Frame(font_mid)
        margin_row.pack(anchor="w")
        ttk.Spinbox(margin_row, from_=0, to=200, textvariable=margin_var, width=8).pack(side="left")
        ttk.Label(margin_row, text="px").pack(side="left", padx=(4, 0))
        add_font_fields(font_bottom, "footer")
        add_color_row(font_bottom, "footer_text_disabled", "押す前の文字色")
        menu_font_inner = section_box(font_tab, "メニュー", fill="x")
        add_color_row(menu_font_inner, "menu_text", "文字色")

        button_top, button_mid, button_bottom = region_columns(button_tab)
        add_button_fields(button_top, "header", "header_button", "header_button_text")
        add_button_fields(button_mid, "", "button", "button_text")
        add_button_fields(button_bottom, "footer", "footer_button", "footer_button_text")
        add_color_row(button_bottom, "footer_button_disabled", "押す前の色")
        border_inner = section_box(button_tab, "縁取り", fill="x", pady=(0, 8))
        border_on = tk.BooleanVar(value=int(self.design.get("border_enabled") or 0) != 0)
        border_w_var = tk.StringVar(value=str(int(self.design.get("border_width") or 1)))
        border_row = ttk.Frame(border_inner)
        border_row.pack(fill="x")
        ttk.Checkbutton(border_row, text="縁取りを付ける", variable=border_on).pack(side="left")
        ttk.Label(border_row, text="縁取りの太さ").pack(side="left", padx=(16, 4))
        ttk.Spinbox(border_row, from_=0, to=8, textvariable=border_w_var, width=8).pack(side="left")
        border_hex = color_value("border")
        color_vars["border"] = tk.StringVar(value=border_hex)
        ttk.Label(border_row, text="縁取りの色").pack(side="left", padx=(16, 4))
        border_swatch = tk.Frame(
            border_row, width=28, height=18, bg=border_hex, highlightthickness=1, highlightbackground="#888"
        )
        border_swatch.pack(side="left")
        border_swatch.pack_propagate(False)
        swatches["border"] = border_swatch
        ttk.Button(border_row, text="色を選ぶ", command=lambda k="border": pick_color(k)).pack(side="left", padx=(8, 0))

        next_box = ttk.LabelFrame(button_tab, text="次へボタンの大きさ", style="Dialog.TLabelframe")
        next_box.pack(fill="x")
        next_inner = ttk.Frame(next_box)
        next_inner.pack(fill="x", padx=12, pady=(8, 14))
        mode_var = tk.StringVar(
            value=(get(self.config, "display", "next_size_mode", NEXT_SIZE_MODE_DEFAULT) or NEXT_SIZE_MODE_DEFAULT).lower()
        )
        if mode_var.get() not in ("percent", "px"):
            mode_var.set("percent")
        ttk.Radiobutton(
            next_inner,
            text="窓全体に対する割合 (%)",
            variable=mode_var,
            value="percent",
            command=lambda: sync_size_mode(),
        ).grid(row=0, column=0, sticky="w")
        ttk.Radiobutton(
            next_inner,
            text="絶対サイズ (px)",
            variable=mode_var,
            value="px",
            command=lambda: sync_size_mode(),
        ).grid(row=1, column=0, sticky="w")
        width_var = tk.StringVar(
            value=str(get(self.config, "display", "next_width", str(NEXT_WIDTH_DEFAULT)) or str(NEXT_WIDTH_DEFAULT))
        )
        height_var = tk.StringVar(
            value=str(get(self.config, "display", "next_height", str(NEXT_HEIGHT_DEFAULT)) or str(NEXT_HEIGHT_DEFAULT))
        )
        size_fields = ttk.Frame(next_inner)
        size_fields.grid(row=0, column=1, rowspan=2, sticky="w", padx=(16, 0))
        width_cap = ttk.Label(size_fields, text="横")
        height_cap = ttk.Label(size_fields, text="縦")
        width_cap.pack(side="left")
        ttk.Entry(size_fields, textvariable=width_var, width=10).pack(side="left", padx=(4, 12))
        height_cap.pack(side="left")
        ttk.Entry(size_fields, textvariable=height_var, width=10).pack(side="left", padx=(4, 0))
        align_var = tk.StringVar(value=(get(self.config, "display", "next_align", "center") or "center").lower())
        if align_var.get() not in ("left", "center", "right"):
            align_var.set("center")
        ttk.Label(next_inner, text="位置").grid(row=2, column=0, sticky="w", pady=(8, 2))
        align_row = ttk.Frame(next_inner)
        align_row.grid(row=3, column=0, columnspan=2, sticky="w")
        ttk.Radiobutton(align_row, text="左", variable=align_var, value="left").pack(side="left", padx=(0, 12))
        ttk.Radiobutton(align_row, text="中央", variable=align_var, value="center").pack(side="left", padx=(0, 12))
        ttk.Radiobutton(align_row, text="右", variable=align_var, value="right").pack(side="left")

        def sync_size_mode() -> None:
            if mode_var.get() == "percent":
                width_cap.config(text="横（窓の幅に対する %）")
                height_cap.config(text="縦（窓の高さに対する %）")
            else:
                width_cap.config(text="横 (px)")
                height_cap.config(text="縦 (px)")

        def reset_design() -> None:
            apply_design_form(
                default_design(),
                {
                    "fullscreen": "0",
                    "next_size_mode": NEXT_SIZE_MODE_DEFAULT,
                    "next_width": str(NEXT_WIDTH_DEFAULT),
                    "next_height": str(NEXT_HEIGHT_DEFAULT),
                },
            )

        window_inner = section_box(window_tab, "表示", fill="x", pady=(0, 8))
        fullscreen_var = tk.BooleanVar(value=get_int(self.config, "display", "fullscreen", 0) != 0)
        hide_close_var = tk.BooleanVar(value=get_int(self.config, "display", "hide_close", 0) != 0)
        display_row = ttk.Frame(window_inner)
        display_row.pack(fill="x")
        ttk.Button(display_row, text="デザインのリセット", command=reset_design).pack(side="right")
        ttk.Checkbutton(
            display_row, text="全画面モード", variable=fullscreen_var, command=lambda: sync_close()
        ).pack(side="left")
        hide_check = ttk.Checkbutton(display_row, text="×ボタンを隠す（全画面のとき）", variable=hide_close_var)
        hide_check.pack(side="left", padx=(16, 0))
        page_var = tk.BooleanVar(value=get_int(self.config, "display", "show_page", 1) != 0)
        ttk.Checkbutton(display_row, text="ページ数を表示", variable=page_var).pack(side="left", padx=(16, 0))

        def sync_close() -> None:
            hide_check.state(["!disabled"] if fullscreen_var.get() else ["disabled"])

        bg_top, bg_mid, bg_bottom = region_columns(window_tab)
        add_color_row(bg_top, "header_background", "背景色")
        add_color_row(bg_mid, "background", "背景色")
        add_color_row(bg_bottom, "footer_background", "背景色")
        extra_inner = section_box(window_tab, "その他", fill="x")
        extra_colors = ttk.Frame(extra_inner)
        extra_colors.pack(fill="x")
        for index, (key, label) in enumerate(
            (
                ("menu", "メニューの色（ドロップダウンを開いたとき）"),
                ("textbox", "テキストボックスの背景色"),
                ("gadget", "ガジェットの色"),
            )
        ):
            col = ttk.Frame(extra_colors)
            col.pack(side="left", fill="both", expand=True, padx=(0 if index == 0 else 8, 0))
            add_color_row(col, key, label)
        title_block = ttk.Frame(extra_inner)
        title_block.pack(fill="x", pady=(8, 0))
        ttk.Label(title_block, text="タイトルのフォント").pack(anchor="w", pady=(0, 2))
        title_fields = ttk.Frame(title_block)
        title_fields.pack(anchor="w")
        title_combo = ttk.Combobox(
            title_fields, textvariable=title_font_var, values=families, width=22, state="readonly"
        )
        title_combo.pack(side="left")
        if title_font_var.get() not in families:
            title_combo.config(state="normal")
        ttk.Label(title_fields, text="文字の大きさ").pack(side="left", padx=(16, 4))
        ttk.Spinbox(title_fields, from_=8, to=72, textvariable=title_size_var, width=8).pack(side="left")
        title_hex = color_value("title_text")
        color_vars["title_text"] = tk.StringVar(value=title_hex)
        ttk.Label(title_fields, text="文字色").pack(side="left", padx=(16, 4))
        title_swatch = tk.Frame(
            title_fields, width=28, height=18, bg=title_hex, highlightthickness=1, highlightbackground="#888"
        )
        title_swatch.pack(side="left")
        title_swatch.pack_propagate(False)
        swatches["title_text"] = title_swatch
        ttk.Button(title_fields, text="色を選ぶ", command=lambda k="title_text": pick_color(k)).pack(side="left", padx=(8, 0))

        def _sync_font_combo(combo: ttk.Combobox, value: str) -> None:
            combo.config(state="normal" if value not in families else "readonly")

        def apply_design_form(data: dict[str, Any], display: dict[str, str] | None = None) -> None:
            font_vars[""].set(str(data.get("font_family") or DEFAULTS["font_family"]))
            font_vars["header"].set(str(data.get("header_font_family") or data.get("font_family") or DEFAULTS["font_family"]))
            font_vars["footer"].set(str(data.get("footer_font_family") or data.get("font_family") or DEFAULTS["font_family"]))
            size_vars[""].set(str(int(data.get("font_size") or DEFAULTS["font_size"])))
            size_vars["header"].set(str(int(data.get("header_font_size") or data.get("font_size") or DEFAULTS["font_size"])))
            size_vars["footer"].set(str(int(data.get("footer_font_size") or data.get("font_size") or DEFAULTS["footer_font_size"])))
            bold_vars[""].set(int(data.get("font_bold") or 0) != 0)
            bold_vars["header"].set(int(data.get("header_font_bold") or 0) != 0)
            bold_vars["footer"].set(int(data.get("footer_font_bold") or 0) != 0)
            button_size_vars[""].set(str(int(data.get("button_font_size") or data.get("font_size") or DEFAULTS["button_font_size"])))
            button_size_vars["header"].set(
                str(int(data.get("header_button_font_size") or data.get("header_font_size") or DEFAULTS["header_button_font_size"]))
            )
            button_size_vars["footer"].set(
                str(int(data.get("footer_button_font_size") or data.get("footer_font_size") or DEFAULTS["footer_button_font_size"]))
            )
            title_font_var.set(str(data.get("title_font_family") or data.get("font_family") or DEFAULTS["title_font_family"]))
            title_size_var.set(str(int(data.get("title_font_size") or DEFAULTS["title_font_size"])))
            margin_var.set(str(int(data.get("content_margin", DEFAULTS["content_margin"]))))
            border_on.set(int(data.get("border_enabled") or 0) != 0)
            border_w_var.set(str(int(data.get("border_width") or 1)))
            for key, var in color_vars.items():
                value = str(data.get(key) or "#f0f0f0")
                var.set(value)
                swatches[key].configure(bg=value)
            for region, combo in font_combos.items():
                _sync_font_combo(combo, font_vars[region].get())
            _sync_font_combo(title_combo, title_font_var.get())
            if display:
                if "fullscreen" in display:
                    fullscreen_var.set(str(display.get("fullscreen") or "0") != "0")
                    sync_close()
                if "hide_close" in display:
                    hide_close_var.set(str(display.get("hide_close") or "0") != "0")
                if "show_page" in display:
                    page_var.set(str(display.get("show_page") or "0") != "0")
                mode = str(display.get("next_size_mode") or "").lower()
                if mode in ("percent", "px"):
                    mode_var.set(mode)
                if "next_width" in display:
                    width_var.set(str(display["next_width"]))
                if "next_height" in display:
                    height_var.set(str(display["next_height"]))
                align = str(display.get("next_align") or "").lower()
                if align in ("left", "center", "right"):
                    align_var.set(align)
            sync_size_mode()

        def collect_form() -> tuple[dict[str, Any], dict[str, Any]] | None:
            try:
                next_width = float(width_var.get().strip())
                next_height = float(height_var.get().strip())
            except ValueError:
                messagebox.showerror("表示のカスタマイズ", "次へボタンの大きさは数値で指定してください。", parent=win)
                return None
            try:
                font_size = int(float(size_vars[""].get().strip()))
                header_font_size = int(float(size_vars["header"].get().strip()))
                footer_font_size = int(float(size_vars["footer"].get().strip()))
                button_font_size = int(float(button_size_vars[""].get().strip()))
                header_button_font_size = int(float(button_size_vars["header"].get().strip()))
                footer_button_font_size = int(float(button_size_vars["footer"].get().strip()))
                title_font_size = int(float(title_size_var.get().strip()))
                border_width = int(float(border_w_var.get().strip()))
                content_margin = int(float(margin_var.get().strip()))
            except ValueError:
                messagebox.showerror("表示のカスタマイズ", "フォントと縁取りの大きさは数値で指定してください。", parent=win)
                return None
            payload = {
                "font_family": font_vars[""].get().strip() or str(DEFAULTS["font_family"]),
                "font_size": font_size,
                "font_bold": 1 if bold_vars[""].get() else 0,
                "header_font_family": font_vars["header"].get().strip() or font_vars[""].get().strip() or str(DEFAULTS["font_family"]),
                "header_font_size": header_font_size,
                "header_font_bold": 1 if bold_vars["header"].get() else 0,
                "footer_font_family": font_vars["footer"].get().strip() or font_vars[""].get().strip() or str(DEFAULTS["font_family"]),
                "footer_font_size": footer_font_size,
                "footer_font_bold": 1 if bold_vars["footer"].get() else 0,
                "button_font_size": button_font_size,
                "header_button_font_size": header_button_font_size,
                "footer_button_font_size": footer_button_font_size,
                "title_font_family": title_font_var.get().strip() or font_vars[""].get().strip() or str(DEFAULTS["title_font_family"]),
                "title_font_size": title_font_size,
                "content_margin": content_margin,
                "border_enabled": 1 if border_on.get() else 0,
                "border_width": border_width,
            }
            for key, var in color_vars.items():
                payload[key] = var.get()
                payload[f"{key}_rgb"] = hex_to_rgb(var.get(), (240, 240, 240))
            display = {
                "fullscreen": "1" if fullscreen_var.get() else "0",
                "hide_close": "1" if hide_close_var.get() else "0",
                "show_page": "1" if page_var.get() else "0",
                "next_size_mode": mode_var.get(),
                "next_width": str(next_width),
                "next_height": str(next_height),
                "next_align": align_var.get(),
            }
            return payload, display

        def save_and_close() -> None:
            collected = collect_form()
            if collected is None:
                return
            payload, display = collected
            if not self.config.has_section("display"):
                self.config.add_section("display")
            for key, value in display.items():
                self.config.set("display", key, str(value))
            save_config(self.config)
            self.design = save_design(payload)
            win.destroy()
            self._next_align = None
            self._next_size = None
            self._apply_design()
            self._apply_window_mode()
            self._layout_next_button()
            self._show_current_ui()

        preset_var = tk.StringVar(value="")
        file_inner = section_box(file_tab, "ファイル", fill="x")
        ttk.Label(file_inner, text="今の内容をファイルに保存できます。読み出すとこの窓の内容が変わります。").pack(anchor="w")
        file_line = ttk.Frame(file_inner)
        file_line.pack(fill="x", pady=(8, 0))
        ttk.Entry(file_line, textvariable=preset_var).pack(side="left", fill="x", expand=True)
        ttk.Button(file_line, text="参照", command=lambda: browse_preset()).pack(side="left", padx=(8, 0))
        file_actions = ttk.Frame(file_inner)
        file_actions.pack(anchor="w", pady=(12, 0))
        ttk.Button(file_actions, text="保存", command=lambda: save_preset()).pack(side="left")
        ttk.Button(file_actions, text="読み出し", command=lambda: load_preset()).pack(side="left", padx=(8, 0))

        def browse_preset() -> None:
            raw = preset_var.get().strip()
            current = resolve_user_path(raw, "design.ini") if raw else ROOT / "design.ini"
            start = current.parent if current.parent.exists() else ROOT
            chosen = filedialog.askopenfilename(
                parent=win,
                title="表示のカスタマイズ",
                initialdir=str(start),
                filetypes=[("設定", "*.ini"), ("すべて", "*.*")],
            )
            if chosen:
                preset_var.set(path_for_config(Path(chosen)))

        def save_preset() -> None:
            collected = collect_form()
            if collected is None:
                return
            payload, display = collected
            raw = preset_var.get().strip()
            if raw:
                path = resolve_user_path(raw, raw)
            else:
                chosen = filedialog.asksaveasfilename(
                    parent=win,
                    title="表示のカスタマイズを保存",
                    initialdir=str(ROOT),
                    defaultextension=".ini",
                    initialfile="design.ini",
                    filetypes=[("設定", "*.ini"), ("すべて", "*.*")],
                )
                if not chosen:
                    return
                path = Path(chosen)
                preset_var.set(path_for_config(path))
            if path.suffix.lower() != ".ini":
                path = path.with_suffix(".ini")
                preset_var.set(path_for_config(path))
            try:
                path.parent.mkdir(parents=True, exist_ok=True)
                save_design_preset(payload, path, display)
            except OSError as exc:
                messagebox.showerror("表示のカスタマイズ", f"保存できませんでした。\n{exc}", parent=win)
                return
            messagebox.showinfo("表示のカスタマイズ", f"保存しました。\n{path.name}", parent=win)

        def load_preset() -> None:
            raw = preset_var.get().strip()
            path = resolve_user_path(raw, raw) if raw else None
            if path is None or not path.is_file():
                chosen = filedialog.askopenfilename(
                    parent=win,
                    title="表示のカスタマイズを読み出す",
                    initialdir=str(ROOT),
                    filetypes=[("設定", "*.ini"), ("すべて", "*.*")],
                )
                if not chosen:
                    return
                path = Path(chosen)
                preset_var.set(path_for_config(path))
            if not path.is_file():
                messagebox.showerror("表示のカスタマイズ", f"ファイルが見つかりません。\n{path}", parent=win)
                return
            try:
                data, display = load_design_preset(path)
            except OSError as exc:
                messagebox.showerror("表示のカスタマイズ", f"読み出せませんでした。\n{exc}", parent=win)
                return
            apply_design_form(data, display)
            messagebox.showinfo("表示のカスタマイズ", "読み出しました。決定で反映します。", parent=win)

        buttons = ttk.Frame(frame)
        buttons.pack(side="bottom", fill="x", pady=(8, 0))
        ttk.Button(buttons, text="決定", command=save_and_close).pack(side="right")
        ttk.Button(buttons, text="キャンセル", command=win.destroy).pack(side="right", padx=(0, 8))
        notebook.pack(fill="both", expand=True, pady=(0, 8))
        sync_close()
        sync_size_mode()
        self._ensure_dialog_styles()
        self._skin_dialog(win)
        win.geometry(f"+{self.root.winfo_rootx() + 40}+{self.root.winfo_rooty() + 40}")

    def _cui_log_name(self) -> str:
        raw = get(self.config, "cui", "file", "")
        if raw:
            return Path(raw).name
        try:
            return resolve_cui_path().name
        except FileNotFoundError:
            return ""

    def _make_logger(self, directory: Path, cui_name: str) -> SessionLogger:
        return SessionLogger(
            directory,
            cui_name,
            backup_directory=resolve_backup_log_directory(self.config),
        )

    def _configured_cui_path(self) -> Path | None:
        raw = get(self.config, "cui", "file", "")
        if not raw:
            return None
        path = resolve_user_path(raw, raw)
        return path if path.is_file() else None

    def _cui_needs_prompt(self) -> bool:
        if get_int(self.config, "cui", "remember", 1) == 0:
            return True
        return self._configured_cui_path() is None

    def _after_main_shown(self, tries: int = 0) -> None:
        try:
            self.root.update_idletasks()
        except tk.TclError:
            return
        ready = bool(self.root.winfo_ismapped()) and self.root.winfo_width() > 1
        if not ready and tries < 40:
            self.root.after(50, lambda: self._after_main_shown(tries + 1))
            return
        if self._startup_cui_started:
            return
        self._startup_cui_started = True
        self._startup_cui()

    def _startup_cui(self) -> None:
        if self._cui_needs_prompt():
            self.root.after(50, self._open_cui_load_dialog)
            return
        path = self._configured_cui_path()
        if path is None:
            self.root.after(50, self._open_cui_load_dialog)
            return
        try:
            self._apply_cui(path)
        except Exception as exc:
            messagebox.showerror("CUIを読み込む", f"CUI を読み込めませんでした。\n{exc}", parent=self.root)
            self.root.after(50, self._open_cui_load_dialog)

    def _apply_cui(self, path: Path) -> None:
        entries = load_cui(path)
        if not self.config.has_section("cui"):
            self.config.add_section("cui")
        self.config.set("cui", "file", path_for_config(path))
        self.entries = entries
        preload_uis(entry.ui_name for entry in self.entries)
        self.index = 0
        self._last_logged = None
        self.logger = self._make_logger(resolve_log_directory(self.config), path.name)
        self._show_current_ui()

    def _place_cui_load_dialog(self, win: tk.Toplevel) -> None:
        try:
            if not win.winfo_exists():
                return
        except tk.TclError:
            return
        try:
            self.root.update_idletasks()
        except tk.TclError:
            return
        width, height = 520, 140
        root_w = max(width, self.root.winfo_width())
        root_h = max(height, self.root.winfo_height())
        x = self.root.winfo_rootx() + max(40, (root_w - width) // 2)
        y = self.root.winfo_rooty() + max(40, (root_h - height) // 2)
        win.geometry(f"{width}x{height}+{x}+{y}")
        win.deiconify()
        win.lift()
        try:
            win.attributes("-topmost", True)
        except tk.TclError:
            pass
        win.focus_force()
        try:
            win.grab_set()
        except tk.TclError:
            pass

    def _open_cui_load_dialog(self) -> None:
        existing = self._cui_load_win
        if existing is not None:
            try:
                if existing.winfo_exists():
                    self._place_cui_load_dialog(existing)
                    return
            except tk.TclError:
                pass
            self._cui_load_win = None

        win = tk.Toplevel(self.root)
        self._cui_load_win = win
        win.withdraw()
        win.title("CUIを読み込む")
        win.transient(self.root)
        win.resizable(True, False)
        win.geometry("520x140")
        self._ensure_dialog_styles()

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        current = self._configured_cui_path()
        initial = path_for_config(current) if current is not None else (get(self.config, "cui", "file", "") or "cui.xlsx")
        cui_var = tk.StringVar(value=initial)

        ttk.Label(frame, text="読み込む CUI").grid(row=0, column=0, sticky="w")
        ttk.Entry(frame, textvariable=cui_var).grid(row=0, column=1, sticky="ew", padx=(8, 8))
        ttk.Button(frame, text="参照", command=lambda: browse()).grid(row=0, column=2, sticky="e")

        def browse() -> None:
            chosen_dir = resolve_user_path(cui_var.get(), "cui.xlsx")
            chosen = filedialog.askopenfilename(
                parent=win,
                title="CUIを読み込む",
                initialdir=str(chosen_dir.parent if chosen_dir.parent.exists() else ROOT),
                filetypes=[
                    ("Excel", "*.xlsx"),
                    ("Excel", "*.xlsm"),
                    ("Excel", "*.xls"),
                    ("CSV", "*.csv"),
                    ("すべて", "*.*"),
                ],
            )
            if chosen:
                cui_var.set(path_for_config(Path(chosen)))

        def close_popup() -> None:
            try:
                win.attributes("-topmost", False)
            except tk.TclError:
                pass
            try:
                win.grab_release()
            except tk.TclError:
                pass
            if self._cui_load_win is win:
                self._cui_load_win = None
            try:
                win.destroy()
            except tk.TclError:
                pass
            if not self.entries:
                self._show_current_ui()

        def apply() -> None:
            path = resolve_user_path(cui_var.get(), "cui.xlsx")
            if not path.is_file():
                messagebox.showerror("CUIを読み込む", f"CUIファイルが見つかりません。\n{path}", parent=win)
                return
            if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls", ".csv"}:
                messagebox.showerror("CUIを読み込む", "CUI は Excel（.xlsx）または CSV を指定してください。", parent=win)
                return
            try:
                load_cui(path)
            except Exception as exc:
                messagebox.showerror("CUIを読み込む", f"CUI を読み込めませんでした。\n{exc}", parent=win)
                return
            if not self.config.has_section("cui"):
                self.config.add_section("cui")
            self.config.set("cui", "file", path_for_config(path))
            save_config(self.config)
            self._apply_cui(path)
            close_popup()

        buttons = ttk.Frame(frame)
        buttons.grid(row=1, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="開く", command=apply).pack(side="right")
        ttk.Button(buttons, text="キャンセル", command=close_popup).pack(side="right", padx=(0, 8))
        win.protocol("WM_DELETE_WINDOW", close_popup)
        self._skin_dialog(win)
        self.root.after(50, lambda: self._place_cui_load_dialog(win))

    def _open_settings_dialog(self) -> None:
        win = tk.Toplevel(self.root)
        win.title("設定")
        win.transient(self.root)
        win.resizable(True, False)
        win.grab_set()
        win.geometry("560x400")

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        try:
            cui_now = path_for_config(resolve_cui_path())
        except FileNotFoundError:
            cui_now = get(self.config, "cui", "file", "cui.xlsx") or "cui.xlsx"
        log_now = get(self.config, "log", "directory", "logs") or "logs"
        backup_now = get(self.config, "log", "backup", "")
        cui_var = tk.StringVar(value=cui_now)
        log_var = tk.StringVar(value=log_now)
        backup_var = tk.StringVar(value=backup_now)
        remember_var = tk.BooleanVar(value=get_int(self.config, "cui", "remember", 1) != 0)
        sound_var = tk.BooleanVar(value=get_int(self.config, "sound", "next", 1) != 0)
        back_off_var = tk.BooleanVar(value=get_int(self.config, "nav", "back", 1) == 0)
        comment_var = tk.BooleanVar(value=get_int(self.config, "nav", "comment", 0) != 0)

        ttk.Label(frame, text="読み込む CUI").grid(row=0, column=0, sticky="w", pady=(0, 8))
        ttk.Entry(frame, textvariable=cui_var).grid(row=0, column=1, sticky="ew", padx=(8, 8), pady=(0, 8))
        ttk.Button(frame, text="参照", command=lambda: browse_cui()).grid(row=0, column=2, sticky="e", pady=(0, 8))
        ttk.Checkbutton(frame, text="前回の情報を残す", variable=remember_var).grid(
            row=1, column=0, columnspan=3, sticky="w", pady=(0, 12)
        )

        ttk.Label(frame, text="ログの保存先").grid(row=2, column=0, sticky="w", pady=(0, 8))
        ttk.Entry(frame, textvariable=log_var).grid(row=2, column=1, sticky="ew", padx=(8, 8), pady=(0, 8))
        ttk.Button(frame, text="参照", command=lambda: browse_log()).grid(row=2, column=2, sticky="e", pady=(0, 8))

        ttk.Label(frame, text="予備のログ保存先").grid(row=3, column=0, sticky="w")
        ttk.Entry(frame, textvariable=backup_var).grid(row=3, column=1, sticky="ew", padx=(8, 8))
        ttk.Button(frame, text="参照", command=lambda: browse_backup()).grid(row=3, column=2, sticky="e")

        ttk.Checkbutton(frame, text="次へを押したときに音を鳴らす", variable=sound_var).grid(
            row=4, column=0, columnspan=3, sticky="w", pady=(12, 0)
        )
        ttk.Checkbutton(frame, text="戻るを使わない", variable=back_off_var).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=(8, 0)
        )
        ttk.Checkbutton(frame, text="メニューの次へでコメントを書く", variable=comment_var).grid(
            row=6, column=0, columnspan=3, sticky="w", pady=(8, 0)
        )

        def browse_cui() -> None:
            current = resolve_user_path(cui_var.get(), "cui.xlsx")
            chosen = filedialog.askopenfilename(
                parent=win,
                title="読み込む CUI",
                initialdir=str(current.parent if current.parent.exists() else ROOT),
                filetypes=[
                    ("Excel", "*.xlsx"),
                    ("Excel", "*.xlsm"),
                    ("Excel", "*.xls"),
                    ("CSV", "*.csv"),
                    ("すべて", "*.*"),
                ],
            )
            if chosen:
                cui_var.set(path_for_config(Path(chosen)))

        def browse_log() -> None:
            current = resolve_user_path(log_var.get(), "logs")
            chosen = filedialog.askdirectory(
                parent=win,
                title="ログの保存先",
                initialdir=str(current if current.exists() else ROOT),
            )
            if chosen:
                log_var.set(path_for_config(Path(chosen)))

        def browse_backup() -> None:
            raw = backup_var.get().strip()
            current = resolve_user_path(raw, raw) if raw else ROOT
            chosen = filedialog.askdirectory(
                parent=win,
                title="予備のログ保存先",
                initialdir=str(current if current.exists() else ROOT),
            )
            if chosen:
                backup_var.set(path_for_config(Path(chosen)))

        def save_and_close() -> None:
            cui_path = resolve_user_path(cui_var.get(), "cui.xlsx")
            if not cui_path.exists():
                messagebox.showerror("設定", f"CUIファイルが見つかりません。\n{cui_path}", parent=win)
                return
            if cui_path.suffix.lower() not in {".xlsx", ".xlsm", ".xls", ".csv"}:
                messagebox.showerror("設定", "CUI は Excel（.xlsx）または CSV を指定してください。", parent=win)
                return
            log_path = resolve_user_path(log_var.get(), "logs")
            try:
                log_path.mkdir(parents=True, exist_ok=True)
            except OSError as exc:
                messagebox.showerror("設定", f"ログの保存先を作れませんでした。\n{exc}", parent=win)
                return
            backup_raw = backup_var.get().strip()
            backup_path = resolve_user_path(backup_raw, backup_raw) if backup_raw else None

            old_cui = resolve_user_path(get(self.config, "cui", "file", "cui.xlsx") or "cui.xlsx", "cui.xlsx")
            old_log = resolve_log_directory(self.config)
            reload_cui = old_cui.resolve() != cui_path.resolve()
            restart_log = old_log.resolve() != log_path.resolve() or reload_cui
            new_entries = None
            if reload_cui:
                try:
                    new_entries = load_cui(cui_path)
                except Exception as exc:
                    messagebox.showerror("設定", f"CUI を読み込めませんでした。\n{exc}", parent=win)
                    return

            if not self.config.has_section("cui"):
                self.config.add_section("cui")
            if not self.config.has_section("log"):
                self.config.add_section("log")
            if not self.config.has_section("sound"):
                self.config.add_section("sound")
            if not self.config.has_section("nav"):
                self.config.add_section("nav")
            self.config.set("cui", "file", path_for_config(cui_path))
            self.config.set("cui", "remember", "1" if remember_var.get() else "0")
            self.config.set("log", "directory", path_for_config(log_path))
            self.config.set("log", "backup", path_for_config(backup_path) if backup_path is not None else "")
            self.config.set("sound", "next", "1" if sound_var.get() else "0")
            self.config.set("nav", "back", "0" if back_off_var.get() else "1")
            self.config.set("nav", "comment", "1" if comment_var.get() else "0")
            save_config(self.config)

            if new_entries is not None:
                self.entries = new_entries
                preload_uis(entry.ui_name for entry in self.entries)
                self.index = 0
                self._last_logged = None
                self._show_current_ui()
            if restart_log:
                self.logger = self._make_logger(log_path, self._cui_log_name())
            else:
                logger = getattr(self, "logger", None)
                if logger is not None:
                    logger.set_backup(backup_path)
            self._refresh_back_menu()
            win.destroy()

        buttons = ttk.Frame(frame)
        buttons.grid(row=7, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="決定", command=save_and_close).pack(side="right")
        ttk.Button(buttons, text="キャンセル", command=win.destroy).pack(side="right", padx=(0, 8))
        self._ensure_dialog_styles()
        self._skin_dialog(win)
        win.geometry(f"+{self.root.winfo_rootx() + 80}+{self.root.winfo_rooty() + 80}")

    def _open_parse_dialog(self) -> None:
        limit = max(channel_max(self.config), RAW_CHANNEL)
        win = tk.Toplevel(self.root)
        win.title("通信データの加工")
        win.transient(self.root)
        win.resizable(True, False)
        win.grab_set()

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)

        box = ttk.LabelFrame(frame, text="データ番号ごと")
        box.pack(fill="x")
        headers = ("番号", "加工", "開始", "終了", "何番目", "送り方")
        for col, text in enumerate(headers):
            ttk.Label(box, text=text).grid(row=0, column=col, padx=8, pady=(6, 4), sticky="w")

        mode_labels = [PARSE_MODE_LABELS[name] for name in PARSE_MODES]
        output_labels = [PARSE_OUTPUT_LABELS[name] for name in PARSE_OUTPUTS]
        mode_from_label = {label: name for name, label in PARSE_MODE_LABELS.items()}
        output_from_label = {label: name for name, label in PARSE_OUTPUT_LABELS.items()}

        mode_vars: list[tk.StringVar] = []
        start_vars: list[tk.StringVar] = []
        end_vars: list[tk.StringVar] = []
        csv_vars: list[tk.StringVar] = []
        output_vars: list[tk.StringVar] = []
        start_fields: list[ttk.Entry] = []
        end_fields: list[ttk.Entry] = []
        csv_fields: list[ttk.Entry] = []

        def sync_row(index: int) -> None:
            mode = mode_from_label.get(mode_vars[index].get(), PARSE_MODE_NONE)
            start_fields[index].config(state="normal" if mode == PARSE_MODE_SLICE else "disabled")
            end_fields[index].config(state="normal" if mode == PARSE_MODE_SLICE else "disabled")
            csv_fields[index].config(state="normal" if mode == PARSE_MODE_CSV else "disabled")

        for number in range(1, limit + 1):
            spec = load_parse_spec(self.config, number)
            row = number
            ttk.Label(box, text=channel_label(number)).grid(row=row, column=0, padx=8, pady=2, sticky="w")
            mode_var = tk.StringVar(value=PARSE_MODE_LABELS[spec["mode"]])
            mode_combo = ttk.Combobox(box, textvariable=mode_var, values=mode_labels, state="readonly", width=14)
            mode_combo.grid(row=row, column=1, padx=(0, 8), pady=2)
            start_var = tk.StringVar(value=str(spec["start"]))
            start_field = ttk.Entry(box, textvariable=start_var, width=6)
            start_field.grid(row=row, column=2, padx=(0, 8), pady=2)
            end_var = tk.StringVar(value=str(spec["end"]))
            end_field = ttk.Entry(box, textvariable=end_var, width=6)
            end_field.grid(row=row, column=3, padx=(0, 8), pady=2)
            csv_var = tk.StringVar(value=str(spec["csv"]))
            csv_field = ttk.Entry(box, textvariable=csv_var, width=6)
            csv_field.grid(row=row, column=4, padx=(0, 8), pady=2)
            output_var = tk.StringVar(value=PARSE_OUTPUT_LABELS[spec["output"]])
            output_combo = ttk.Combobox(box, textvariable=output_var, values=output_labels, state="readonly", width=8)
            output_combo.grid(row=row, column=5, padx=(0, 8), pady=2)
            mode_vars.append(mode_var)
            start_vars.append(start_var)
            end_vars.append(end_var)
            csv_vars.append(csv_var)
            output_vars.append(output_var)
            start_fields.append(start_field)
            end_fields.append(end_field)
            csv_fields.append(csv_field)
            mode_combo.bind("<<ComboboxSelected>>", lambda _e, i=number - 1: sync_row(i))
            sync_row(number - 1)

        ttk.Label(
            frame,
            text="文字列の一部は1文字目から数えた開始〜終了です。カンマ区切りは左から何番目かです。",
        ).pack(anchor="w", pady=(8, 0))

        def save_and_close() -> None:
            if not self.config.has_section("parse"):
                self.config.add_section("parse")
            for index, number in enumerate(range(1, limit + 1)):
                key = _parse_channel_key(number)
                mode = mode_from_label.get(mode_vars[index].get(), PARSE_MODE_NONE)
                output = output_from_label.get(output_vars[index].get(), PARSE_OUTPUT_TEXT)
                try:
                    start = int(float(start_vars[index].get().strip()))
                    end = int(float(end_vars[index].get().strip()))
                    csv_n = int(float(csv_vars[index].get().strip()))
                except ValueError:
                    messagebox.showerror("通信データの加工", "開始・終了・何番目は数値で指定してください。", parent=win)
                    return
                if start < 1 or end < 1 or csv_n < 1:
                    messagebox.showerror("通信データの加工", "開始・終了・何番目は1以上にしてください。", parent=win)
                    return
                if end < start:
                    start, end = end, start
                self.config.set("parse", f"{key}_mode", mode)
                self.config.set("parse", f"{key}_output", output)
                self.config.set("parse", f"{key}_start", str(start))
                self.config.set("parse", f"{key}_end", str(end))
                self.config.set("parse", f"{key}_csv", str(csv_n))
            save_config(self.config)
            win.destroy()

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(12, 0))
        ttk.Button(buttons, text="保存", command=save_and_close).pack(side="right")
        ttk.Button(buttons, text="キャンセル", command=win.destroy).pack(side="right", padx=(0, 8))
        win.update_idletasks()
        self._use_default_dialog_look(win)
        win.geometry(f"+{self.root.winfo_rootx() + 80}+{self.root.winfo_rooty() + 80}")

    def _open_com_dialog(self) -> None:
        limit = max(channel_max(self.config), RAW_CHANNEL)
        saved = load_channel_ports(self.config)
        open_ports = self.channels.assigned_ports()
        current = dict(saved)
        current.update(open_ports)

        win = tk.Toplevel(self.root)
        win.title("通信")
        win.transient(self.root)
        win.resizable(False, False)
        win.grab_set()

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)

        use_box = ttk.LabelFrame(frame, text="COM")
        use_box.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        use_var = tk.BooleanVar(value=load_com_enabled(self.config))

        ports_box = ttk.LabelFrame(frame, text="ポート")
        ports_box.grid(row=1, column=0, sticky="ew")

        ttk.Label(ports_box, text="番号").grid(row=0, column=0, padx=(8, 8), sticky="w", pady=(6, 2))
        ttk.Label(ports_box, text="COMポート").grid(row=0, column=1, padx=(0, 8), sticky="w", pady=(6, 2))
        ttk.Label(ports_box, text="ボーレート").grid(row=0, column=2, padx=(0, 8), sticky="w", pady=(6, 2))
        ttk.Label(ports_box, text="ランプ").grid(row=0, column=3, padx=(0, 8), sticky="w", pady=(6, 2))
        ttk.Label(ports_box, text="状態").grid(row=0, column=4, sticky="w", pady=(6, 2), padx=(0, 8))
        ttk.Label(ports_box, text="文字コード").grid(row=0, column=5, sticky="w", pady=(6, 2), padx=(0, 8))
        ttk.Label(ports_box, text="区切り長さ").grid(row=0, column=6, sticky="w", pady=(6, 2), padx=(0, 8))
        ttk.Label(ports_box, text="待ち時間（秒）").grid(row=0, column=7, sticky="w", pady=(6, 2), padx=(0, 8))
        ttk.Label(ports_box, text="再試行").grid(row=0, column=8, sticky="w", pady=(6, 2), padx=(0, 8))

        combos: list[ttk.Combobox] = []
        status_labels: list[ttk.Label] = []
        vars_: list[tk.StringVar] = []
        lamp_vars: list[tk.BooleanVar] = []
        lamp_checks: list[ttk.Checkbutton] = []
        baud_vars: list[tk.StringVar] = []
        baud_combos: list[ttk.Combobox] = []
        enc_vars: list[tk.StringVar] = []
        enc_combos: list[ttk.Combobox] = []
        wait_vars: list[tk.StringVar] = []
        wait_entries: list[ttk.Entry] = []
        retry_vars: list[tk.StringVar] = []
        retry_entries: list[ttk.Entry] = []
        lamp_ports = set(load_lamp_ports(self.config))
        saved_bauds = load_port_baudrates(self.config)
        saved_encs = load_port_encodings(self.config)
        saved_waits = load_port_timeouts(self.config)
        saved_retries = load_port_retries(self.config)
        default_baud = get_int(self.config, "serial", "baudrate", 9600)
        default_wait = clamp_recv_wait(get(self.config, "serial", "timeout", str(DEFAULT_RECV_WAIT)))
        default_retries = clamp_recv_retries(get(self.config, "serial", "retries", str(DEFAULT_RECV_RETRIES)))
        last_ports: list[str] = []
        raw_length_var = tk.StringVar(value=str(load_raw_length(self.config)))
        raw_length_combo: ttk.Combobox | None = None

        def choices_for(selected: str) -> list[str]:
            ports = list_com_ports()
            for port in list(open_ports.values()) + list(saved.values()):
                if port and port not in ports:
                    ports.append(port)
            if selected and selected not in ports and selected != UNUSED_PORT:
                ports.append(selected)
            return [UNUSED_PORT] + ports

        def refresh_status() -> None:
            live = self.channels.assigned_ports()
            for number, label in enumerate(status_labels, start=1):
                port = vars_[number - 1].get().strip()
                if port and live.get(number) == port:
                    label.config(text="接続中")
                elif port and port != UNUSED_PORT:
                    label.config(text="")
                else:
                    label.config(text="未使用")

        def sync_row_enabled() -> None:
            enabled = use_var.get()
            for index, check in enumerate(lamp_checks):
                port = vars_[index].get().strip()
                usable = enabled and bool(port) and port != UNUSED_PORT
                check.state(["!disabled"] if usable else ["disabled"])
                baud_combos[index].config(state="normal" if usable else "disabled")
                enc_combos[index].config(state="readonly" if usable else "disabled")
                wait_entries[index].config(state="normal" if usable else "disabled")
                retry_entries[index].config(state="normal" if usable else "disabled")
                if not usable and not enabled:
                    continue
                if not usable:
                    lamp_vars[index].set(False)
            if raw_length_combo is not None:
                raw_length_combo.config(state="normal" if enabled else "disabled")

        def set_ports_enabled(enabled: bool) -> None:
            combo_state = "readonly" if enabled else "disabled"
            for combo in combos:
                combo.config(state=combo_state)
            reload_btn.state(["!disabled"] if enabled else ["disabled"])
            sync_row_enabled()

        def on_use_toggle() -> None:
            set_ports_enabled(use_var.get())

        def on_combo_selected(_event: object | None = None) -> None:
            refresh_status()
            for index, var in enumerate(vars_):
                port = var.get().strip() or UNUSED_PORT
                if index < len(last_ports) and port == last_ports[index]:
                    continue
                if index < len(last_ports):
                    last_ports[index] = port
                if port and port != UNUSED_PORT:
                    baud_vars[index].set(str(baud_for_port(port, saved_bauds, default_baud)))
                    enc_vars[index].set(encoding_for_port(port, saved_encs))
                    wait_vars[index].set(f"{timeout_for_port(port, saved_waits, default_wait):g}")
                    retry_vars[index].set(str(retries_for_port(port, saved_retries, default_retries)))
            sync_row_enabled()

        def refresh_ports() -> None:
            for number, combo in enumerate(combos, start=1):
                selected = vars_[number - 1].get().strip() or UNUSED_PORT
                values = choices_for(selected)
                combo["values"] = values
                if selected not in values:
                    selected = UNUSED_PORT
                vars_[number - 1].set(selected)
            on_combo_selected()

        for number in range(1, limit + 1):
            ttk.Label(ports_box, text=channel_label(number)).grid(
                row=number, column=0, sticky="w", pady=2, padx=(8, 8)
            )
            initial = current.get(number, "") or UNUSED_PORT
            var = tk.StringVar(value=initial)
            combo = ttk.Combobox(ports_box, textvariable=var, state="readonly", width=14)
            combo.grid(row=number, column=1, sticky="ew", pady=2, padx=(0, 8))
            baud_value = str(default_baud)
            enc_value = DEFAULT_ENCODING
            wait_value = f"{default_wait:g}"
            retry_value = str(default_retries)
            if initial and initial != UNUSED_PORT:
                baud_value = str(baud_for_port(initial, saved_bauds, default_baud))
                enc_value = encoding_for_port(initial, saved_encs)
                wait_value = f"{timeout_for_port(initial, saved_waits, default_wait):g}"
                retry_value = str(retries_for_port(initial, saved_retries, default_retries))
            baud_var = tk.StringVar(value=baud_value)
            baud_combo = ttk.Combobox(ports_box, textvariable=baud_var, values=BAUD_CHOICES, width=10)
            baud_combo.grid(row=number, column=2, sticky="w", pady=2, padx=(0, 8))
            lamp_var = tk.BooleanVar(value=bool(initial) and initial != UNUSED_PORT and initial in lamp_ports)
            check = ttk.Checkbutton(ports_box, variable=lamp_var)
            check.grid(row=number, column=3, padx=(0, 8))
            status = ttk.Label(ports_box, text="")
            status.grid(row=number, column=4, sticky="w", padx=(0, 8))
            enc_var = tk.StringVar(value=enc_value)
            enc_combo = ttk.Combobox(
                ports_box,
                textvariable=enc_var,
                values=ENCODING_CHOICES,
                state="readonly",
                width=10,
            )
            enc_combo.grid(row=number, column=5, sticky="w", padx=(0, 8))
            if number == RAW_CHANNEL:
                raw_length_combo = ttk.Combobox(
                    ports_box,
                    textvariable=raw_length_var,
                    values=RAW_LENGTH_CHOICES,
                    width=8,
                )
                raw_length_combo.grid(row=number, column=6, sticky="w", padx=(0, 8))
            else:
                ttk.Label(ports_box, text="—").grid(row=number, column=6, sticky="w", padx=(0, 8))
            wait_var = tk.StringVar(value=wait_value)
            wait_field = ttk.Entry(ports_box, textvariable=wait_var, width=8)
            wait_field.grid(row=number, column=7, sticky="w", pady=2, padx=(0, 8))
            retry_var = tk.StringVar(value=retry_value)
            retry_field = ttk.Entry(ports_box, textvariable=retry_var, width=6)
            retry_field.grid(row=number, column=8, sticky="w", pady=2, padx=(0, 8))
            vars_.append(var)
            combos.append(combo)
            baud_vars.append(baud_var)
            baud_combos.append(baud_combo)
            enc_vars.append(enc_var)
            enc_combos.append(enc_combo)
            wait_vars.append(wait_var)
            wait_entries.append(wait_field)
            retry_vars.append(retry_var)
            retry_entries.append(retry_field)
            lamp_vars.append(lamp_var)
            lamp_checks.append(check)
            status_labels.append(status)
            last_ports.append(initial)
            combo.bind("<<ComboboxSelected>>", on_combo_selected)

        ttk.Label(
            ports_box,
            text=f"{RAW_CHANNEL_NAME}は改行なしデータ用です。区切り長さ（バイト）ずつ読みます。来たデータはすぐ使えます。待ち時間 0 は待たずに次を見ます（リアルタイム向き）。",
        ).grid(row=limit + 1, column=0, columnspan=9, sticky="w", padx=8, pady=(8, 6))

        ttk.Checkbutton(use_box, text="使う", variable=use_var, command=on_use_toggle).pack(
            anchor="w", padx=8, pady=6
        )

        buttons = ttk.Frame(frame)
        buttons.grid(row=2, column=0, sticky="e", pady=(12, 0))
        reload_btn = ttk.Button(buttons, text="再読み込み", command=refresh_ports)
        reload_btn.pack(side="left", padx=(0, 8))

        def save_and_close() -> None:
            mapping: dict[int, str] = {}
            used: dict[str, int] = {}
            lamps: list[str] = []
            baudrates: dict[str, int] = {}
            encodings: dict[str, str] = {}
            timeouts: dict[str, float] = {}
            retry_counts: dict[str, int] = {}
            for number, var in enumerate(vars_, start=1):
                port = var.get().strip()
                if not port or port == UNUSED_PORT:
                    continue
                if port in used:
                    messagebox.showerror(
                        "通信設定",
                        f"{port} が {channel_label(used[port])} と {channel_label(number)} に重複しています。",
                        parent=win,
                    )
                    return
                rate = parse_baud(baud_vars[number - 1].get(), 0)
                if rate <= 0:
                    messagebox.showerror(
                        "通信設定",
                        f"{port} のボーレートが数値ではありません。",
                        parent=win,
                    )
                    return
                label = encoding_label(enc_vars[number - 1].get())
                if label not in ENCODING_CHOICES:
                    messagebox.showerror(
                        "通信設定",
                        f"{port} の文字コードが不正です。",
                        parent=win,
                    )
                    return
                try:
                    wait_raw = float(wait_vars[number - 1].get().strip())
                except (TypeError, ValueError):
                    messagebox.showerror(
                        "通信設定",
                        f"{port} の待ち時間は数値で指定してください。",
                        parent=win,
                    )
                    return
                if wait_raw < 0:
                    messagebox.showerror(
                        "通信設定",
                        f"{port} の待ち時間は0以上にしてください。",
                        parent=win,
                    )
                    return
                try:
                    retry_raw = int(float(retry_vars[number - 1].get().strip()))
                except (TypeError, ValueError):
                    messagebox.showerror(
                        "通信設定",
                        f"{port} の再試行回数は整数で指定してください。",
                        parent=win,
                    )
                    return
                if retry_raw < 0:
                    messagebox.showerror(
                        "通信設定",
                        f"{port} の再試行回数は0以上にしてください。",
                        parent=win,
                    )
                    return
                used[port] = number
                mapping[number] = port
                baudrates[port] = rate
                encodings[port] = label
                timeouts[port] = clamp_recv_wait(wait_raw)
                retry_counts[port] = clamp_recv_retries(retry_raw)
                if lamp_vars[number - 1].get():
                    lamps.append(port)
            try:
                raw_length = int(float(raw_length_var.get().strip()))
            except (TypeError, ValueError):
                raw_length = 0
            if raw_length < 1:
                messagebox.showerror(
                    "通信設定",
                    f"{RAW_CHANNEL_NAME}の区切り長さは1以上の数値で指定してください。",
                    parent=win,
                )
                return
            use_com = bool(use_var.get())
            self.config = save_channel_ports(
                mapping,
                limit,
                lamps,
                use_com=use_com,
                baudrates=baudrates,
                encodings=encodings,
                timeouts=timeouts,
                retries=retry_counts,
                raw_length=clamp_raw_length(raw_length),
                parser=self.config,
            )
            was_connected = self.channels.connected
            win.destroy()
            self._apply_com_usage()
            if use_com and was_connected:
                self._try_connect(mapping)

        ttk.Button(buttons, text="保存", command=save_and_close).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="キャンセル", command=win.destroy).pack(side="left")
        refresh_ports()
        on_use_toggle()
        win.update_idletasks()
        self._use_default_dialog_look(win)
        win.geometry(f"+{self.root.winfo_rootx() + 80}+{self.root.winfo_rooty() + 80}")

    def _toggle_connection(self) -> None:
        if not load_com_enabled(self.config):
            return
        if self.channels.connected:
            self.channels.disconnect()
            self.connect_button.config(text="接続")
            self._set_lamp(LAMP_OFF)
            self._refresh_condition()
            return

        ports = load_channel_ports(self.config)
        if not ports:
            self._set_lamp(LAMP_ERROR)
            messagebox.showerror("接続エラー", "COM ポートが設定されていません。メニューの「通信」から設定してください。")
            return
        self._try_connect(ports)

    def _try_connect(self, ports: dict[int, str]) -> None:
        if not load_com_enabled(self.config):
            self.channels.disconnect()
            self.connect_button.config(text="接続")
            self._set_lamp(LAMP_OFF)
            self._refresh_condition()
            return
        if not ports:
            self.channels.disconnect()
            self.connect_button.config(text="接続")
            self._set_lamp(LAMP_OFF)
            self._refresh_condition()
            return

        self._set_lamp(LAMP_WAIT)
        self.root.update_idletasks()
        errors = self.channels.connect(
            ports,
            serial_options(self.config),
            load_port_baudrates(self.config),
            encodings=load_port_encodings(self.config),
            timeouts=load_port_timeouts(self.config),
            retries=load_port_retries(self.config),
            raw_length=load_raw_length(self.config),
        )
        if not self.channels.connected:
            self._set_lamp(LAMP_ERROR)
            self.connect_button.config(text="接続")
            detail = "\n".join(errors.values()) if errors else "接続できませんでした。"
            messagebox.showerror("接続エラー", detail)
            self._refresh_condition()
            return

        self.connect_button.config(text="切断")
        self._set_lamp(LAMP_ON)
        if errors:
            messagebox.showwarning("一部接続できませんでした", "\n".join(errors.values()))
        self._refresh_condition()

    def _on_serial_data(self, channel: int, payload: dict[str, Any]) -> None:
        self.root.after(0, lambda ch=channel, data=payload: self._handle_serial_data(ch, data))

    def _handle_serial_data(self, channel: int, payload: dict[str, Any]) -> None:
        if self.current_ui is not None:
            self.current_ui.on_data(channel, self.get_data(channel))
            if channel == 1:
                self.current_ui.on_com_data(payload)
        self._refresh_condition()

    def _current_log_name(self) -> str:
        return self.current_entry.ui_name if self.current_entry is not None else ""

    def _current_log_title(self) -> str:
        if self.current_entry is None:
            return ""
        return str(self.current_entry.title or "")

    def write_log_from_ui(self, content: Any) -> None:
        if isinstance(content, (list, tuple)) and not isinstance(content, (str, bytes)):
            payload: dict[str, Any] | list[Any] = list(content)
        elif isinstance(content, dict):
            payload = {str(key): value for key, value in content.items()}
        else:
            payload = {"log": content}
        if not payload:
            payload = {"log": ""}
        self._emit_log(self._current_log_name(), payload)

    def _judge_now(self) -> str | None:
        if self.current_ui is None:
            return None
        return judge_label(self.current_ui.is_condition_met())

    def _current_log_headings(self) -> list[str]:
        ui = getattr(self, "current_ui", None)
        if ui is None:
            return []
        return normalize_headings(ui.log_headings())

    def _emit_log(self, name: str, values: Any, title: str = "", judge: bool = True) -> None:
        heading = title or self._current_log_title()
        fallback = self._judge_now() if judge else None
        labels = self._current_log_headings()
        if isinstance(values, (list, tuple)) and not isinstance(values, (str, bytes)):
            self.logger.write(name, list(values), title=heading, headings=labels)
            if fallback is not None:
                self.logger.write(name, {JUDGE_KEY: fallback}, title=heading)
            return
        if isinstance(values, dict):
            self.logger.write(name, with_judge_last(values, fallback), title=heading)
            return
        payload: dict[str, Any] = {"log": values}
        if fallback is not None:
            payload[JUDGE_KEY] = fallback
        self.logger.write(name, payload, title=heading)

    def get_data(self, number: int | str) -> Any:
        channel = resolve_channel(number)
        if channel < 1:
            return None
        payload = self.channels.payload(channel)
        if not payload:
            return None
        raw = payload.get("raw")
        if raw is None:
            return None
        return process_received(str(raw), load_parse_spec(self.config, channel))

    def get_raw(self, number: int | str) -> Any:
        return self.channels.raw(number)

    def send_data(self, number: int | str, text: str, flush_input: bool = False) -> bool:
        return self.channels.send(number, str(text or ""), flush_input=bool(flush_input))

    def set_read_end(self, number: int | str, eol: str) -> None:
        self.channels.set_eol(number, str(eol or ""))

    def get_all_data(self) -> dict[int, Any]:
        numbers = set(self.channels.get_all_data().keys())
        return {number: self.get_data(number) for number in sorted(numbers)}

    def done_from_ui(self, value: Any) -> None:
        name = self.current_entry.ui_name if self.current_entry is not None else ""
        self._emit_log(name, done_log_values(value))
        self._refresh_condition()

    def get_com_data(self) -> dict[str, Any] | None:
        return self.channels.payload(1)

    def is_com_connected(self) -> bool:
        return self.channels.connected

    def _clear_content(self) -> None:
        if self.current_ui is not None:
            self.current_ui.destroy()
            self.current_ui = None
        for child in self.content.winfo_children():
            child.destroy()

    def _page_text(self) -> str:
        total = len(self.entries)
        if total <= 0:
            return "0 / 0"
        current = min(self.index + 1, total)
        return f"{current} / {total}"

    def _refresh_page_label(self) -> None:
        label = getattr(self, "page_label", None)
        if label is None:
            return
        show = get_int(self.config, "display", "show_page", 1) != 0
        if not show:
            label.pack_forget()
            return
        label.config(text=self._page_text())
        if not str(label.winfo_manager()):
            label.pack(side="left", padx=(0, 12), before=self.menu_button)

    def _show_current_ui(self) -> None:
        self._clear_content()
        self._refresh_page_label()
        self._refresh_back_menu()
        if not self.entries:
            self._set_next_enabled(False)
            self._refresh_next_caption()
            return
        if self.index >= len(self.entries):
            ttk.Label(self.content, text="すべての UI を表示しました。").grid(sticky="n")
            self.current_entry = None
            self._set_next_enabled(False)
            self._refresh_next_caption()
            return

        entry = self.entries[self.index]
        self.current_entry = entry
        holder = ttk.Frame(self.content)
        holder.grid(row=0, column=0, sticky="nsew")
        holder.columnconfigure(0, weight=1)
        holder.rowconfigure(0, weight=1)

        context = UIContext(
            ui_name=entry.ui_name,
            title=entry.title,
            variables=dict(entry.variables),
            texts=list(entry.texts),
            args=[list(row) for row in entry.args],
            get_data=self.get_data,
            get_all_data=self.get_all_data,
            get_raw=self.get_raw,
            send_data=self.send_data,
            set_read_end=self.set_read_end,
            log_done=self.done_from_ui,
            write_log=self.write_log_from_ui,
            get_com_data=self.get_com_data,
            is_com_connected=self.is_com_connected,
            notify_condition=self._refresh_condition,
            design=dict(self.design),
        )
        try:
            self.current_ui = create_ui(entry.ui_name, holder, context)
        except Exception as exc:
            self.current_ui = None
            for child in holder.winfo_children():
                child.destroy()
            if isinstance(exc, (IndexError, KeyError, ValueError)):
                text = f"{entry.ui_name}\n設定値がCUIから取得できませんでした。"
            else:
                text = f"UI '{entry.ui_name}' を読み込めませんでした。\n{exc}"
            ttk.Label(holder, text=text, justify="center").grid(row=0, column=0)
        self._last_logged = None
        self._refresh_condition()
        self._refresh_next_caption()

    def _condition_met(self) -> bool:
        if self.current_ui is None:
            return False
        return self.current_ui.is_condition_met()

    def _set_next_enabled(self, enabled: bool) -> None:
        self._next_enabled = bool(enabled)
        self._apply_next_button_colors()

    def _play_next_sound(self) -> None:
        if get_int(self.config, "sound", "next", 1) == 0:
            return
        data = self._next_sound_bytes()
        if not data:
            return
        threading.Thread(target=self._emit_wav, args=(data,), daemon=True).start()

    def _next_sound_bytes(self) -> bytes:
        cached = getattr(self, "_next_sound_wav", None)
        if isinstance(cached, bytes):
            return cached
        data = b""
        path = ROOT / "sound1.wav"
        try:
            if path.is_file():
                data = path.read_bytes()
        except OSError:
            data = b""
        self._next_sound_wav = data
        return data

    def _emit_wav(self, data: bytes) -> None:
        try:
            import winsound

            winsound.PlaySound(data, winsound.SND_MEMORY | winsound.SND_NODEFAULT)
        except Exception:
            return

    def _apply_next_button_colors(self) -> None:
        button = getattr(self, "next_button", None)
        if button is None:
            return
        d = self.design
        footer_fg = str(d.get("footer_button_text") or d.get("footer_text") or d.get("text") or "#000000")
        footer_button = str(d.get("footer_button") or d.get("button") or "#f0f0f0")
        enabled = getattr(self, "_next_enabled", False)
        if enabled:
            bg = footer_button
            fg = footer_fg
            active = self._tone_hex(bg, -25)
        else:
            bg = str(d.get("footer_button_disabled") or footer_button)
            fg = str(d.get("footer_text_disabled") or footer_fg)
            active = bg
        # Windows では state=disabled だと背景色が効かないので、色はここで出し、押下は _go_next で止める。
        button.configure(
            state="normal",
            bg=bg,
            fg=fg,
            activebackground=active,
            activeforeground=fg,
            disabledforeground=fg,
            takefocus=1 if enabled else 0,
        )

    def _can_advance(self) -> bool:
        if self.current_entry is None:
            return False
        if self.index >= len(self.entries):
            return False
        return self._condition_met()

    def _refresh_condition(self) -> None:
        self._set_next_enabled(self._can_advance())

    def _log_current(self, force: bool = False) -> None:
        if self.current_entry is None or self.current_ui is None:
            return
        values = self.current_ui.get_received_values()
        if not force and values == self._last_logged:
            return
        self._emit_log(self.current_entry.ui_name, values)
        if isinstance(values, list):
            self._last_logged = list(values)
        elif isinstance(values, dict):
            self._last_logged = dict(values)
        else:
            self._last_logged = None

    def _footer_label(self) -> str:
        ui = getattr(self, "current_ui", None)
        if ui is None:
            return "次へ"
        text = str(ui.next_label() or "").strip()
        return text or "次へ"

    def _footer_exits(self) -> bool:
        ui = getattr(self, "current_ui", None)
        if ui is None:
            return False
        return bool(ui.next_exits())

    def _refresh_next_caption(self) -> None:
        button = getattr(self, "next_button", None)
        if button is None:
            return
        button.configure(text=self._footer_label())

    def _back_enabled(self) -> bool:
        return get_int(self.config, "nav", "back", 1) != 0

    def _comment_on_next(self) -> bool:
        return get_int(self.config, "nav", "comment", 0) == 1

    def _go_next(self) -> None:
        if not self._can_advance():
            return
        self._advance_next()

    def _go_next_from_menu(self) -> None:
        if self.current_entry is None:
            return
        if self._comment_on_next():
            self._open_next_comment_dialog()
            return
        self._advance_next(judge_true=True)

    def _advance_next(self, *, judge_true: bool = False) -> None:
        self._play_next_sound()
        if judge_true:
            self._log_forced_next()
        else:
            self._log_current(force=True)
        if self._footer_exits():
            self._on_exit()
            return
        self.index += 1
        self._show_current_ui()

    def _log_forced_next(self, comment: str = "") -> None:
        ui_name = self.current_entry.ui_name if self.current_entry is not None else ""
        values = self.current_ui.get_received_values() if self.current_ui is not None else {}
        extra: dict[str, Any] = {}
        if comment:
            extra["次へ"] = comment
        extra[JUDGE_KEY] = "TRUE"
        if isinstance(values, list):
            self._emit_log(ui_name, list(values), judge=False)
            self._emit_log(ui_name, extra)
            return
        payload: dict[str, Any] = dict(values) if isinstance(values, dict) else {}
        payload.update(extra)
        self._emit_log(ui_name, payload)

    def _can_go_back(self) -> bool:
        if not self._back_enabled():
            return False
        return bool(self.entries) and self.index > 0

    def _refresh_back_menu(self) -> None:
        menu = getattr(self, "menu", None)
        if menu is None:
            return
        try:
            menu.entryconfig("戻る", state="normal" if self._can_go_back() else "disabled")
        except tk.TclError:
            pass

    def _go_back(self) -> None:
        if not self._can_go_back():
            return
        self.index -= 1
        self._show_current_ui()

    def _current_log_path(self) -> Path | None:
        logger = getattr(self, "logger", None)
        if logger is None:
            return None
        path = getattr(logger, "path", None)
        return Path(path) if path else None

    def _open_record_dialog(self) -> None:
        win = tk.Toplevel(self.root)
        win.title("記録確認")
        win.transient(self.root)
        win.resizable(True, False)
        win.grab_set()
        win.geometry("560x140")

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        current = self._current_log_path()
        initial = path_for_config(current) if current is not None else ""
        log_var = tk.StringVar(value=initial)

        ttk.Label(frame, text="開くログ").grid(row=0, column=0, sticky="w")
        ttk.Entry(frame, textvariable=log_var).grid(row=0, column=1, sticky="ew", padx=(8, 8))

        def browse() -> None:
            chosen_path = resolve_user_path(log_var.get(), "logs")
            start = chosen_path.parent if chosen_path.suffix else chosen_path
            if not start.exists():
                start = resolve_log_directory(self.config)
            chosen = filedialog.askopenfilename(
                parent=win,
                title="開くログ",
                initialdir=str(start if start.exists() else ROOT),
                initialfile=chosen_path.name if chosen_path.suffix else "",
                filetypes=[("CSV", "*.csv"), ("すべて", "*.*")],
            )
            if chosen:
                log_var.set(path_for_config(Path(chosen)))

        ttk.Button(frame, text="参照", command=browse).grid(row=0, column=2, sticky="e")

        def close_popup() -> None:
            try:
                win.grab_release()
            except tk.TclError:
                pass
            win.destroy()

        def apply() -> None:
            raw = str(log_var.get() or "").strip()
            path = resolve_user_path(raw, raw)
            if not raw or not path.is_file():
                messagebox.showerror("記録確認", f"ログファイルが見つかりません。\n{path}", parent=win)
                return
            close_popup()
            self._show_record_dialog(path)

        buttons = ttk.Frame(frame)
        buttons.grid(row=1, column=0, columnspan=3, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="開く", command=apply).pack(side="right")
        ttk.Button(buttons, text="キャンセル", command=close_popup).pack(side="right", padx=(0, 8))
        win.protocol("WM_DELETE_WINDOW", close_popup)
        self._use_default_dialog_look(win)

    def _show_record_dialog(self, path: Path) -> None:
        win = tk.Toplevel(self.root)
        win.title("記録確認")
        win.transient(self.root)
        win.resizable(True, True)
        win.geometry("960x560")
        win.minsize(640, 360)
        win.grab_set()

        look = self._dialog_look()
        font = look["font"]
        bg = look["bg"]
        fg = look["fg"]
        outer = ttk.Frame(win, padding=12, style="Dialog.TFrame")
        outer.pack(fill="both", expand=True)
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(0, weight=1)

        canvas = tk.Canvas(outer, highlightthickness=0, bg=bg)
        vsb = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        hsb = ttk.Scrollbar(outer, orient="horizontal", command=canvas.xview)
        canvas.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        inner = tk.Frame(canvas, bg=bg)
        canvas.create_window((0, 0), window=inner, anchor="nw")

        def _on_inner(_event=None) -> None:
            canvas.configure(scrollregion=canvas.bbox("all") or (0, 0, 0, 0))

        inner.bind("<Configure>", _on_inner)

        def _wheel(event) -> str:
            delta = int(getattr(event, "delta", 0) or 0)
            if delta:
                canvas.yview_scroll(int(-1 * (delta / 120)), "units")
            return "break"

        canvas.bind("<MouseWheel>", _wheel)
        inner.bind("<MouseWheel>", _wheel)

        family = str(font[0]) if isinstance(font, tuple) and font else "ＭＳ Ｐゴシック"
        try:
            size = int(font[1]) if isinstance(font, tuple) and len(font) >= 2 else 9
        except (TypeError, ValueError):
            size = 9
        header_font = (family, size, "bold")

        def put(row: int, col: int, text: str, header: bool = False) -> None:
            label = tk.Label(
                inner,
                text=text,
                font=header_font if header else font,
                bg=bg,
                fg=fg,
                relief="solid",
                bd=1,
                padx=8,
                pady=4,
                anchor="w",
                justify="left",
            )
            label.grid(row=row, column=col, sticky="nsew")

        records = read_log_records(path)
        assigned = _assign_log_blocks(self.entries, records)
        grid, header_cells, gaps, spacers = _record_view_layout(self.entries, assigned)
        for col in gaps:
            inner.columnconfigure(col, minsize=24)
        for row_index, row in enumerate(grid):
            if row_index in spacers:
                inner.rowconfigure(row_index, minsize=max(16, size + 10))
                continue
            for col, cell in enumerate(row):
                if col in gaps and not str(cell).strip():
                    continue
                put(row_index, col, cell, header=(row_index, col) in header_cells)

        def save_named() -> None:
            log_dir = path.parent if path.parent.exists() else resolve_log_directory(self.config)
            suggested = _record_save_name(self._cui_log_name(), manual=True)
            chosen = filedialog.asksaveasfilename(
                parent=win,
                title="記録を保存",
                initialdir=str(log_dir if log_dir.exists() else ROOT),
                initialfile=suggested,
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv"), ("すべて", "*.*")],
            )
            if not chosen:
                return
            dest = Path(chosen)
            if dest.suffix.lower() != ".csv":
                dest = dest.with_suffix(".csv")
            try:
                if path.is_file() and dest.resolve() == path.resolve():
                    messagebox.showinfo("記録確認", "同じファイルです。", parent=win)
                    return
                _write_record_view(dest, grid)
            except OSError as exc:
                messagebox.showerror("記録確認", f"保存できませんでした。\n{exc}", parent=win)
                return
            messagebox.showinfo("記録確認", f"保存しました。\n{dest.name}", parent=win)

        buttons = ttk.Frame(outer)
        buttons.grid(row=2, column=0, columnspan=2, sticky="e", pady=(8, 0))
        ttk.Button(buttons, text="閉じる", command=win.destroy).pack(side="right")
        ttk.Button(buttons, text="保存", command=save_named).pack(side="right", padx=(0, 8))
        win.protocol("WM_DELETE_WINDOW", win.destroy)
        self._use_default_dialog_look(win)

    def _open_next_comment_dialog(self) -> None:
        if self.current_entry is None:
            return
        win = tk.Toplevel(self.root)
        win.title("次へ")
        win.transient(self.root)
        win.resizable(True, True)
        win.grab_set()
        win.geometry(f"+{self.root.winfo_rootx() + 80}+{self.root.winfo_rooty() + 80}")

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        text = tk.Text(frame, width=40, height=6, wrap="word")
        text.pack(fill="both", expand=True)
        text.focus_set()

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(8, 0))

        def close_popup() -> None:
            win.grab_release()
            win.destroy()

        def confirm() -> None:
            comment = str(text.get("1.0", "end-1c") or "").strip()
            close_popup()
            self._play_next_sound()
            self._log_forced_next(comment)
            if self._footer_exits():
                self._on_exit()
                return
            self.index += 1
            self._show_current_ui()

        ttk.Button(buttons, text="OK", command=confirm).pack(side="right")
        ttk.Button(buttons, text="取り消し", command=close_popup).pack(side="right", padx=(0, 8))
        win.protocol("WM_DELETE_WINDOW", close_popup)
        self._use_default_dialog_look(win)

    def _jump_to_cui(self, index: int) -> None:
        if index < 0 or index >= len(self.entries):
            return
        self.index = index
        self._show_current_ui()

    def _open_ui_list_dialog(self) -> None:
        win = tk.Toplevel(self.root)
        win.title("UI一覧")
        win.transient(self.root)
        win.resizable(True, True)
        win.grab_set()
        win.geometry(f"520x360+{self.root.winfo_rootx() + 80}+{self.root.winfo_rooty() + 80}")

        frame = ttk.Frame(win, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(0, weight=1)

        if not self.entries:
            ttk.Label(frame, text="CUI に UI がありません。").pack(anchor="w")
            ttk.Button(frame, text="閉じる", command=win.destroy).pack(anchor="e", pady=(12, 0))
            win.protocol("WM_DELETE_WINDOW", win.destroy)
            self._use_default_dialog_look(win)
            return

        columns = ("ui_name", "title")
        tree = ttk.Treeview(frame, columns=columns, show="headings", selectmode="browse")
        tree.heading("ui_name", text="UI")
        tree.heading("title", text="タイトル")
        tree.column("ui_name", width=140, stretch=False)
        tree.column("title", width=320, stretch=True)
        scroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")

        for index, entry in enumerate(self.entries):
            title = (entry.title or "").strip() or "（タイトルなし）"
            tree.insert("", "end", iid=str(index), values=(entry.ui_name, title))
        if 0 <= self.index < len(self.entries):
            current = str(self.index)
            tree.selection_set(current)
            tree.focus(current)
            tree.see(current)

        def selected_index() -> int | None:
            chosen = tree.selection()
            if not chosen:
                return None
            try:
                return int(chosen[0])
            except (TypeError, ValueError):
                return None

        def close_popup() -> None:
            win.grab_release()
            win.destroy()

        def open_selected(_event: object | None = None) -> None:
            index = selected_index()
            if index is None:
                return
            close_popup()
            self._jump_to_cui(index)

        tree.bind("<Double-1>", open_selected)
        tree.bind("<Return>", open_selected)

        buttons = ttk.Frame(frame)
        buttons.grid(row=1, column=0, columnspan=2, sticky="e", pady=(8, 0))
        ttk.Button(buttons, text="開く", command=open_selected).pack(side="right")
        ttk.Button(buttons, text="閉じる", command=close_popup).pack(side="right", padx=(0, 8))
        win.protocol("WM_DELETE_WINDOW", close_popup)
        self._use_default_dialog_look(win)

    def _tick_condition(self) -> None:
        self._refresh_condition()
        self._refresh_com_lamps()
        self.root.after(200, self._tick_condition)

    def _save_record_backup(self, grid: list[list[str]], filename: str) -> None:
        backup = resolve_backup_log_directory(self.config)
        if backup is None:
            return
        try:
            if backup.resolve() == resolve_log_directory(self.config).resolve():
                return
            backup.mkdir(parents=True, exist_ok=True)
            dest = backup / filename
            extra = 1
            while dest.exists():
                dest = dest.with_name(f"{dest.stem}_{extra}{dest.suffix}")
                extra += 1
            _write_record_view(dest, grid)
        except OSError:
            return

    def _save_record_on_exit(self) -> None:
        logger = getattr(self, "logger", None)
        records = logger.read_records() if logger is not None else []
        assigned = _assign_log_blocks(self.entries, records)
        grid, _headers, _gaps, _spacers = _record_view_layout(self.entries, assigned)
        log_dir = resolve_log_directory(self.config)
        primary_error: OSError | None = None
        dest_name = _record_save_name(self._cui_log_name())
        try:
            log_dir.mkdir(parents=True, exist_ok=True)
            dest = log_dir / dest_name
            extra = 1
            while dest.exists():
                dest = dest.with_name(f"{dest.stem}_{extra}{dest.suffix}")
                extra += 1
            _write_record_view(dest, grid)
            dest_name = dest.name
        except OSError as exc:
            primary_error = exc
        self._save_record_backup(grid, dest_name)
        if primary_error is not None:
            raise primary_error

    def _on_exit(self) -> None:
        if getattr(self, "_exiting", False):
            return
        self._exiting = True
        try:
            self._save_record_on_exit()
        except OSError as exc:
            messagebox.showerror("終了", f"記録を保存できませんでした。\n{exc}")
        self.channels.disconnect()
        if get_int(self.config, "display", "fullscreen", 0) == 0:
            if not self.config.has_section("window"):
                self.config.add_section("window")
            self.config.set("window", "width", str(max(640, self.root.winfo_width())))
            self.config.set("window", "height", str(max(480, self.root.winfo_height())))
        save_config(self.config)
        self.root.destroy()


def run() -> None:
    root = tk.Tk()
    MainApp(root)
    root.mainloop()


if __name__ == "__main__":
    run()
